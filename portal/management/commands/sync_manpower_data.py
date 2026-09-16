import os
import re
import json
from datetime import datetime, date
import openpyxl
from django.core.management.base import BaseCommand
from django.db import transaction
from portal.models import LabourRecord, RVJVEmployee, HiredOperator, ContractorWorkerPPE

def parse_date(val):
    if not val:
        return None
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    if isinstance(val, str):
        val = val.strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d/%m/%y', '%d-%m-%y'):
            try:
                return datetime.strptime(val, fmt).date()
            except ValueError:
                pass
    return None

def clean_str(val, max_len=None):
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() == 'none' or s.lower() == 'nan':
        return None
    if max_len and len(s) > max_len:
        s = s[:max_len]
    return s

class Command(BaseCommand):
    help = 'Sync Manpower data from Excel files into database models, or export/import via JSON'

    def add_arguments(self, parser):
        parser.add_argument('--export', action='store_true', help='Export manpower tables to JSON')
        parser.add_argument('--import-json', action='store_true', help='Import manpower tables from JSON')
        parser.add_argument('--excel-dir', type=str, default=r'C:\Users\adev9\OneDrive\Desktop\safety_office_data', help='Directory with Excel files')

    def handle(self, *args, **options):
        base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
        json_file = os.path.join(base_dir, 'manpower_data_export.json')

        if options.get('export'):
            self.export_json(json_file)
            return

        if options.get('import_json'):
            self.import_json(json_file)
            return

        excel_dir = options.get('excel_dir')
        self.stdout.write(self.style.NOTICE(f'Loading Excel files from: {excel_dir}'))

        self.import_labour(excel_dir)
        self.import_rvjv(excel_dir)
        self.import_hiring(excel_dir)
        self.import_contractor_ppe(excel_dir)

        self.stdout.write(self.style.SUCCESS(
            f'\nManpower Sync Summary:\n'
            f'  - LabourRecord: {LabourRecord.objects.count()}\n'
            f'  - RVJVEmployee: {RVJVEmployee.objects.count()}\n'
            f'  - HiredOperator: {HiredOperator.objects.count()}\n'
            f'  - ContractorWorkerPPE: {ContractorWorkerPPE.objects.count()}\n'
        ))

    def import_labour(self, excel_dir, specific_file=None):
        path = specific_file or os.path.join(excel_dir, 'Labour man power List.xlsx')
        if not os.path.exists(path):
            self.stdout.write(self.style.WARNING(f'File not found: {path}'))
            return

        self.stdout.write('Importing Labour Records...')
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb['Sheet1']

        objects = []
        seen_ids = set()

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            labour_id = clean_str(row[1], 50)
            if not labour_id or labour_id in seen_ids:
                continue
            seen_ids.add(labour_id)

            name = clean_str(row[2], 200) or 'Unnamed Worker'
            gender = clean_str(row[3], 20)
            cid = clean_str(row[4], 50)
            nationality = clean_str(row[5], 50) or 'Indian'
            if nationality not in ('Bhutanese', 'Indian', 'Other'):
                nationality = 'Other'
            status = clean_str(row[6], 20) or 'Active'
            if status not in ('Active', 'Leave', 'Left', 'Registered', 'Transferred'):
                status = 'Active'
            voter_id = clean_str(row[7], 50)

            objects.append(LabourRecord(
                labour_id=labour_id,
                name=name,
                gender=gender,
                nationality=nationality,
                cid_number=cid,
                voter_id=voter_id,
                status=status,
            ))

        with transaction.atomic():
            LabourRecord.objects.all().delete()
            LabourRecord.objects.bulk_create(objects, batch_size=500)
        self.stdout.write(self.style.SUCCESS(f'  -> Imported {len(objects)} Labour records.'))

    def import_rvjv(self, excel_dir, specific_file=None):
        path = specific_file or os.path.join(excel_dir, 'Man power of rvjv.xlsx')
        if not os.path.exists(path):
            self.stdout.write(self.style.WARNING(f'File not found: {path}'))
            return

        self.stdout.write('Importing RVJV Employees...')
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb['Sheet1']

        objects = []
        seen_ids = set()

        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[1]:
                continue
            emp_id = clean_str(row[1], 50)
            if not emp_id or emp_id in seen_ids:
                continue
            seen_ids.add(emp_id)

            name = clean_str(row[2], 200) or 'Unnamed Employee'
            designation = clean_str(row[3], 200)
            department = clean_str(row[4], 200)
            doj = parse_date(row[5])
            nationality = clean_str(row[6], 50) or 'Bhutanese'
            if nationality not in ('Bhutanese', 'Indian', 'Other'):
                nationality = 'Other'
            wp = clean_str(row[7], 100)
            cid = clean_str(row[8], 50)
            status = clean_str(row[9], 20) or 'Active'
            if status not in ('Active', 'Inactive', 'Left'):
                status = 'Active'
            relieving = parse_date(row[10])
            mobile = clean_str(row[11], 30)

            objects.append(RVJVEmployee(
                emp_id=emp_id,
                name=name,
                designation=designation,
                department=department,
                date_of_joining=doj,
                nationality=nationality,
                work_permit=wp,
                cid_number=cid,
                status=status,
                relieving_date=relieving,
                mobile=mobile,
            ))

        with transaction.atomic():
            RVJVEmployee.objects.all().delete()
            RVJVEmployee.objects.bulk_create(objects, batch_size=500)
        self.stdout.write(self.style.SUCCESS(f'  -> Imported {len(objects)} RVJV employees.'))

    def import_hiring(self, excel_dir, specific_file=None):
        path = os.path.join(excel_dir, 'Hiring Man power List.xlsx')
        if not os.path.exists(path):
            self.stdout.write(self.style.WARNING(f'File not found: {path}'))
            return

        self.stdout.write('Importing Hired Operators...')
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb['Sheet1']

        objects = []
        seen_ids = set()

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            op_id = clean_str(row[1], 50)
            if not op_id or op_id in seen_ids:
                continue
            seen_ids.add(op_id)

            name = clean_str(row[2], 200) or 'Unnamed Operator'
            veh = clean_str(row[3], 50)
            agent = clean_str(row[4], 200)
            cid = clean_str(row[5], 50)
            raw_ht = clean_str(row[6], 50) or 'Driver/Operator'
            if 'mechanic' in raw_ht.lower() or 'supervisor' in raw_ht.lower():
                hire_type = 'Mechanic/Supervisor'
            elif 'driver' in raw_ht.lower() or 'operator' in raw_ht.lower():
                hire_type = 'Driver/Operator'
            else:
                hire_type = 'Other'
            desig = clean_str(row[7], 200)
            nat = clean_str(row[8], 50) or 'Bhutanese'
            if nat not in ('Bhutanese', 'Indian', 'Other'):
                nat = 'Other'
            doj = parse_date(row[9])

            objects.append(HiredOperator(
                operator_id=op_id,
                name=name,
                vehicle_no=veh,
                hire_agent=agent,
                cid_no=cid,
                hire_type=hire_type,
                designation=desig,
                nationality=nat,
                date_of_joining=doj,
            ))

        with transaction.atomic():
            HiredOperator.objects.all().delete()
            HiredOperator.objects.bulk_create(objects, batch_size=500)
        self.stdout.write(self.style.SUCCESS(f'  -> Imported {len(objects)} Hired Operators.'))

    def import_contractor_ppe(self, excel_dir, specific_file=None):
        path = specific_file or os.path.join(excel_dir, 'HIRINGS.xlsx')
        if not os.path.exists(path):
            self.stdout.write(self.style.WARNING(f'File not found: {path}'))
            return

        self.stdout.write('Importing Contractor Worker PPE issuances across all sheets...')
        wb = openpyxl.load_workbook(path, data_only=True)

        objects = []
        seen_keys = set()

        for sname in wb.sheetnames:
            if sname in ('Sheet1', 'Sheet2'):
                continue
            ws = wb[sname]

            h_row = None
            for r in range(1, min(6, ws.max_row+1)):
                vals = [str(ws.cell(r, c).value or '').strip().lower() for c in range(1, min(ws.max_column+1, 10))]
                if any('name' in v for v in vals) and (any('id' in v for v in vals) or any('sl' in v for v in vals)):
                    h_row = r
                    break
            if not h_row:
                continue

            col_map = {}
            for c in range(1, ws.max_column+1):
                raw_h = ws.cell(h_row, c).value
                if raw_h:
                    norm = str(raw_h).strip().lower()
                    col_map[c] = norm

            name_c = None
            id_c = None
            desig_c = None
            ppe_cols = {}

            for c, norm in col_map.items():
                if ('name' in norm) and not name_c and ('contractor' not in norm) and ('hiring' not in norm):
                    name_c = c
                elif ('id' in norm or norm == 'sl no' or norm == 'sl. no' or norm == 'sl. n o' or norm == 'sl') and not id_c:
                    id_c = c
                elif ('degis' in norm or 'desig' in norm or 'degina' in norm) and not desig_c:
                    desig_c = c
                elif 'reflector' in norm:
                    ppe_cols['reflector'] = c
                elif 'safety boot' in norm or 'boot' in norm:
                    ppe_cols['safety_boot'] = c
                elif 'helmet' in norm and 'welding' not in norm:
                    ppe_cols['helmet'] = c
                elif 'glove' in norm and 'welding' not in norm:
                    ppe_cols['gloves'] = c
                elif 'mask' in norm:
                    ppe_cols['face_mask'] = c
                elif 'goggle' in norm and 'black' not in norm:
                    ppe_cols['safety_goggles'] = c
                elif 'ear' in norm:
                    ppe_cols['ear_plug'] = c
                elif 'shoulder' in norm:
                    ppe_cols['shoulder_pads'] = c
                elif 'harness' in norm:
                    ppe_cols['body_harness'] = c
                elif 'raincoat' in norm:
                    ppe_cols['raincoat'] = c

            contractor_name = sname.strip()

            for r in range(h_row+1, ws.max_row+1):
                name_val = ws.cell(r, name_c).value if name_c else None
                name_str = clean_str(name_val, 200)
                if not name_str:
                    continue

                w_id = clean_str(ws.cell(r, id_c).value, 50) if id_c else None
                if not w_id:
                    w_id = f'ROW-{r}'

                desig_val = clean_str(ws.cell(r, desig_c).value, 200) if desig_c else None

                key = (w_id, contractor_name)
                if key in seen_keys:
                    w_id = f'{w_id}_{r}'
                    key = (w_id, contractor_name)
                seen_keys.add(key)

                ppe_data = {}
                for f_name, c_idx in ppe_cols.items():
                    val = ws.cell(r, c_idx).value
                    ppe_data[f_name] = clean_str(val, 500)

                objects.append(ContractorWorkerPPE(
                    worker_id=w_id,
                    name=name_str,
                    contractor_name=contractor_name,
                    designation=desig_val,
                    reflector=ppe_data.get('reflector'),
                    safety_boot=ppe_data.get('safety_boot'),
                    helmet=ppe_data.get('helmet'),
                    gloves=ppe_data.get('gloves'),
                    face_mask=ppe_data.get('face_mask'),
                    safety_goggles=ppe_data.get('safety_goggles'),
                    ear_plug=ppe_data.get('ear_plug'),
                    shoulder_pads=ppe_data.get('shoulder_pads'),
                    body_harness=ppe_data.get('body_harness'),
                    raincoat=ppe_data.get('raincoat'),
                ))

        with transaction.atomic():
            ContractorWorkerPPE.objects.all().delete()
            ContractorWorkerPPE.objects.bulk_create(objects, batch_size=500)
        self.stdout.write(self.style.SUCCESS(f'  -> Imported {len(objects)} Contractor Worker PPE records from {len(wb.sheetnames)} sheets.'))

    def export_json(self, json_file):
        self.stdout.write('Exporting all manpower records to JSON...')
        data = {
            'labour': list(LabourRecord.objects.values('labour_id', 'name', 'gender', 'nationality', 'cid_number', 'voter_id', 'status')),
            'rvjv': [
                {
                    'emp_id': r.emp_id,
                    'name': r.name,
                    'designation': r.designation,
                    'department': r.department,
                    'nationality': r.nationality,
                    'date_of_joining': r.date_of_joining.strftime('%Y-%m-%d') if r.date_of_joining else None,
                    'work_permit': r.work_permit,
                    'cid_number': r.cid_number,
                    'status': r.status,
                    'relieving_date': r.relieving_date.strftime('%Y-%m-%d') if r.relieving_date else None,
                    'mobile': r.mobile,
                }
                for r in RVJVEmployee.objects.all()
            ],
            'hired': [
                {
                    'operator_id': r.operator_id,
                    'name': r.name,
                    'vehicle_no': r.vehicle_no,
                    'hire_agent': r.hire_agent,
                    'cid_no': r.cid_no,
                    'hire_type': r.hire_type,
                    'designation': r.designation,
                    'nationality': r.nationality,
                    'date_of_joining': r.date_of_joining.strftime('%Y-%m-%d') if r.date_of_joining else None,
                }
                for r in HiredOperator.objects.all()
            ],
            'contractor_ppe': list(ContractorWorkerPPE.objects.values(
                'worker_id', 'name', 'contractor_name', 'designation',
                'reflector', 'safety_boot', 'helmet', 'gloves', 'face_mask',
                'safety_goggles', 'ear_plug', 'shoulder_pads', 'body_harness', 'raincoat'
            ))
        }

        with open(json_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        self.stdout.write(self.style.SUCCESS(
            f'Exported to {json_file}:\n'
            f'  Labour: {len(data["labour"])}\n'
            f'  RVJV: {len(data["rvjv"])}\n'
            f'  Hired: {len(data["hired"])}\n'
            f'  Contractor PPE: {len(data["contractor_ppe"])}\n'
        ))

    def import_json(self, json_file):
        if not os.path.exists(json_file):
            self.stdout.write(self.style.ERROR(f'JSON export file not found: {json_file}'))
            return

        with open(json_file, 'r', encoding='utf-8') as f:
            data = json.load(f)

        with transaction.atomic():
            LabourRecord.objects.all().delete()
            labour_objs = [LabourRecord(**item) for item in data.get('labour', [])]
            LabourRecord.objects.bulk_create(labour_objs, batch_size=500)

            RVJVEmployee.objects.all().delete()
            rvjv_objs = []
            for item in data.get('rvjv', []):
                doj = datetime.strptime(item['date_of_joining'], '%Y-%m-%d').date() if item.get('date_of_joining') else None
                rel = datetime.strptime(item['relieving_date'], '%Y-%m-%d').date() if item.get('relieving_date') else None
                rvjv_objs.append(RVJVEmployee(
                    emp_id=item['emp_id'],
                    name=item['name'],
                    designation=item.get('designation'),
                    department=item.get('department'),
                    nationality=item.get('nationality', 'Bhutanese'),
                    date_of_joining=doj,
                    work_permit=item.get('work_permit'),
                    cid_number=item.get('cid_number'),
                    status=item.get('status', 'Active'),
                    relieving_date=rel,
                    mobile=item.get('mobile'),
                ))
            RVJVEmployee.objects.bulk_create(rvjv_objs, batch_size=500)

            HiredOperator.objects.all().delete()
            hired_objs = []
            for item in data.get('hired', []):
                doj = datetime.strptime(item['date_of_joining'], '%Y-%m-%d').date() if item.get('date_of_joining') else None
                hired_objs.append(HiredOperator(
                    operator_id=item['operator_id'],
                    name=item['name'],
                    vehicle_no=item.get('vehicle_no'),
                    hire_agent=item.get('hire_agent'),
                    cid_no=item.get('cid_no'),
                    hire_type=item.get('hire_type', 'Driver/Operator'),
                    designation=item.get('designation'),
                    nationality=item.get('nationality', 'Bhutanese'),
                    date_of_joining=doj,
                ))
            HiredOperator.objects.bulk_create(hired_objs, batch_size=500)

            ContractorWorkerPPE.objects.all().delete()
            ppe_objs = [ContractorWorkerPPE(**item) for item in data.get('contractor_ppe', [])]
            ContractorWorkerPPE.objects.bulk_create(ppe_objs, batch_size=500)

        self.stdout.write(self.style.SUCCESS(
            f'Successfully restored from JSON:\n'
            f'  Labour: {LabourRecord.objects.count()}\n'
            f'  RVJV: {RVJVEmployee.objects.count()}\n'
            f'  Hired: {HiredOperator.objects.count()}\n'
            f'  Contractor PPE: {ContractorWorkerPPE.objects.count()}\n'
        ))
