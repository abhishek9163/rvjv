import datetime
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from decimal import Decimal

from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db.models import Sum, Count, Q, F

from .models import (
    User,
    Employee,
    SafetyStoreItem,
    SafetyEquipmentIssue,
    SafetyReplacementAndFine
)


def has_safety_access(user):
    """
    Check whether a user has permission to access the Safety Department module.
    Permitted: Superusers, Managers, or users with 'safety_department' in assigned_modules.
    """
    if not user.is_authenticated:
        return False
    if user.is_superuser or user.system_role == 'MANAGER':
        return True
    assigned = list(getattr(user, 'assigned_modules', []) or [])
    return 'safety_department' in assigned


# ==============================================================================
# MAIN DASHBOARD VIEW
# ==============================================================================

@login_required
def safety_dashboard_view(request):
    if not has_safety_access(request.user):
        messages.error(request, "Access Restricted: You do not have permission to access the Safety Department module.")
        return redirect('dashboard')

    today = timezone.now().date()
    soon_threshold = today + datetime.timedelta(days=30)

    # All store items with creator
    items_qs = SafetyStoreItem.objects.select_related('created_by').all().order_by('category', 'name')
    total_store_types = items_qs.count()
    total_stock_all = sum(i.total_stock for i in items_qs)
    total_available_all = sum(i.available_stock for i in items_qs)
    total_issued_all = sum(i.issued_stock for i in items_qs)
    low_stock_items = [i for i in items_qs if i.is_low_stock]

    # Equipment Issues
    all_issues_qs = SafetyEquipmentIssue.objects.select_related('employee', 'item', 'issued_by').all().order_by('-issue_date', '-created_at')
    active_issues_qs = all_issues_qs.filter(status='ACTIVE')
    active_issues_count = active_issues_qs.count()

    expiring_soon_count = 0
    expired_warranty_count = 0
    for iss in active_issues_qs:
        if iss.warranty_expiry_date:
            if iss.warranty_expiry_date < today:
                expired_warranty_count += 1
            elif iss.warranty_expiry_date <= soon_threshold:
                expiring_soon_count += 1

    # Replacements & Fines
    fines_qs = SafetyReplacementAndFine.objects.select_related('employee', 'item', 'original_issue', 'processed_by').all().order_by('-replacement_date', '-created_at')
    pending_fines = fines_qs.filter(fine_status='PENDING_APPROVAL')
    pending_fines_count = pending_fines.count()
    total_pending_fine_amount = pending_fines.aggregate(s=Sum('fine_amount'))['s'] or Decimal('0.00')
    total_fine_collected = fines_qs.filter(fine_status__in=['APPROVED_SALARY_DEDUCTION', 'PAID_CASH']).aggregate(s=Sum('fine_amount'))['s'] or Decimal('0.00')


    # NOTE: Employee dropdowns now use AJAX search (/safety/api/employees/search/)
    # No need to pass employees queryset to context — avoids rendering thousands of options

    # Categories list - dynamically gathered from existing store items + standard options
    db_categories = list(SafetyStoreItem.objects.exclude(category__isnull=True).exclude(category='').values_list('category', flat=True).distinct())
    standard_categories = [
        'Head Protection',
        'Safety Shoes',
        'Reflectors & Hi-Vis',
        'Hand Protection',
        'Body Harness & Fall Protection',
        'Eye & Face Protection',
        'Hearing Protection',
        'Respiratory Protection',
        'General PPE',
        'Other'
    ]
    all_categories = sorted(list(set(db_categories + standard_categories)))

    reasons = SafetyReplacementAndFine.REASON_CHOICES
    fine_statuses = SafetyReplacementAndFine.FINE_STATUS_CHOICES
    issue_statuses = SafetyEquipmentIssue.STATUS_CHOICES

    departments = sorted(list(set(filter(None, Employee.objects.exclude(department__isnull=True).exclude(department='').values_list('department', flat=True).distinct()))))
    creators = sorted(list(set(filter(None, SafetyStoreItem.objects.exclude(created_by__isnull=True).values_list('created_by__username', flat=True).distinct()))))

    context = {
        'today': today,
        'store_items': items_qs,
        'low_stock_items': low_stock_items,
        'total_store_types': total_store_types,
        'total_stock_all': total_stock_all,
        'total_available_all': total_available_all,
        'total_issued_all': total_issued_all,
        'low_stock_count': len(low_stock_items),
        'active_issues_count': active_issues_count,
        'expiring_soon_count': expiring_soon_count,
        'expired_warranty_count': expired_warranty_count,
        'pending_fines_count': pending_fines_count,
        'total_pending_fine_amount': total_pending_fine_amount,
        'total_fine_collected': total_fine_collected,
        'issues': all_issues_qs[:150],
        'fines': fines_qs[:150],
        'categories': all_categories,
        'departments': departments,
        'creators': creators,
        'reasons': reasons,
        'fine_statuses': fine_statuses,
        'issue_statuses': issue_statuses,
        'active_tab': request.GET.get('tab', 'inventory'),
    }

    return render(request, 'safety_dashboard.html', context)


# ==============================================================================
# STORE INVENTORY MANAGEMENT APIS (CREATE, EDIT, DELETE, RESTOCK)
# ==============================================================================

@login_required
def api_safety_item_create(request):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    try:
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST
        name = data.get('name', '').strip()
        item_code = data.get('item_code', '').strip().upper()
        category = data.get('category', 'OTHER').strip()
        unit = data.get('unit', 'Pairs').strip()
        specification = data.get('specification', '').strip()
        
        try:
            total_stock = int(data.get('total_stock', 0))
            if total_stock < 0: total_stock = 0
        except:
            total_stock = 0

        try:
            warranty_months = int(data.get('warranty_months', 6))
            if warranty_months < 1: warranty_months = 1
        except:
            warranty_months = 6

        try:
            fine_amount = Decimal(str(data.get('fine_amount', '0.00')))
            if fine_amount < 0: fine_amount = Decimal('0.00')
        except:
            fine_amount = Decimal('0.00')

        try:
            minimum_alert_level = int(data.get('minimum_alert_level', 5))
        except:
            minimum_alert_level = 5

        notes = data.get('notes', '').strip()

        # Fallbacks for optional fields
        if not name:
            name = "Safety Equipment"

        if not item_code:
            import uuid
            item_code = f"SAF-{uuid.uuid4().hex[:6].upper()}"
            while SafetyStoreItem.objects.filter(item_code__iexact=item_code).exists():
                item_code = f"SAF-{uuid.uuid4().hex[:6].upper()}"
        elif SafetyStoreItem.objects.filter(item_code__iexact=item_code).exists():
            return JsonResponse({'status': 'ERROR', 'msg': f"Item code '{item_code}' is already registered."})

        item = SafetyStoreItem.objects.create(
            name=name,
            item_code=item_code,
            category=category,
            unit=unit,
            specification=specification,
            total_stock=total_stock,
            available_stock=total_stock,
            warranty_months=warranty_months,
            fine_amount=fine_amount,
            minimum_alert_level=minimum_alert_level,
            created_by=request.user,
            notes=notes
        )

        return JsonResponse({
            'status': 'SUCCESS',
            'msg': f"Equipment '{item.name}' ({item.item_code}) added to safety store with initial stock {item.available_stock} {item.unit}.",
            'item_id': item.id
        })
    except Exception as e:
        return JsonResponse({'status': 'ERROR', 'msg': str(e)})


@login_required
def api_safety_item_edit(request, item_id):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    item = get_object_or_404(SafetyStoreItem, id=item_id)
    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    try:
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST
        name = data.get('name', '').strip()
        item_code = data.get('item_code', '').strip().upper()
        category = data.get('category', item.category).strip()
        unit = data.get('unit', item.unit).strip()
        specification = data.get('specification', '').strip()

        if not name or not item_code:
            return JsonResponse({'status': 'ERROR', 'msg': 'Item Name and Item Code are required.'})

        if SafetyStoreItem.objects.filter(item_code__iexact=item_code).exclude(id=item.id).exists():
            return JsonResponse({'status': 'ERROR', 'msg': f"Item code '{item_code}' already in use by another item."})

        item.name = name
        item.item_code = item_code
        item.category = category
        item.unit = unit
        item.specification = specification

        if 'warranty_months' in data:
            try:
                item.warranty_months = max(1, int(data['warranty_months']))
            except:
                pass

        if 'fine_amount' in data:
            try:
                item.fine_amount = max(Decimal('0.00'), Decimal(str(data['fine_amount'])))
            except:
                pass

        if 'minimum_alert_level' in data:
            try:
                item.minimum_alert_level = max(0, int(data['minimum_alert_level']))
            except:
                pass

        if 'notes' in data:
            item.notes = data['notes']

        if 'available_stock' in data:
            try:
                new_avail = max(0, int(data['available_stock']))
                diff = new_avail - item.available_stock
                item.available_stock = new_avail
                item.total_stock = max(0, item.total_stock + diff)
            except:
                pass

        item.save()
        return JsonResponse({'status': 'SUCCESS', 'msg': f"Equipment '{item.name}' updated successfully."})
    except Exception as e:
        return JsonResponse({'status': 'ERROR', 'msg': str(e)})


@login_required
def api_safety_item_delete(request, item_id):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    item = get_object_or_404(SafetyStoreItem, id=item_id)
    active_count = item.issue_records.filter(status='ACTIVE').count()
    if active_count > 0:
        return JsonResponse({
            'status': 'ERROR',
            'msg': f"Cannot delete '{item.name}' because {active_count} active issue record(s) are currently assigned to workers."
        })

    item_name = item.name
    item.delete()
    return JsonResponse({'status': 'SUCCESS', 'msg': f"Equipment item '{item_name}' removed from store catalog."})


@login_required
def api_safety_item_restock(request, item_id):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    item = get_object_or_404(SafetyStoreItem, id=item_id)
    try:
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST
        qty = int(data.get('quantity', 0))
        if qty <= 0:
            return JsonResponse({'status': 'ERROR', 'msg': 'Restock quantity must be greater than 0.'})

        item.total_stock += qty
        item.available_stock += qty
        item.save()

        return JsonResponse({
            'status': 'SUCCESS',
            'msg': f"Added +{qty} {item.unit} to '{item.name}'. Current Available Stock: {item.available_stock} {item.unit}."
        })
    except Exception as e:
        return JsonResponse({'status': 'ERROR', 'msg': str(e)})


# ==============================================================================
# WORKER EQUIPMENT ISSUANCE APIS (CREATE, EDIT, DELETE, GET WORKER ACTIVE)
# ==============================================================================

@login_required
def api_safety_issue_create(request):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    try:
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST
        emp_id = data.get('employee_id')
        item_id = data.get('item_id')
        quantity = int(data.get('quantity', 1))
        size_specification = data.get('size_specification', '').strip()
        issue_date_str = data.get('issue_date')
        remarks = data.get('remarks', '').strip()

        if not emp_id or not item_id:
            return JsonResponse({'status': 'ERROR', 'msg': 'Both Employee and Equipment are required.'})

        employee = get_object_or_404(Employee, id=emp_id)
        item = get_object_or_404(SafetyStoreItem, id=item_id)

        if item.available_stock < quantity:
            return JsonResponse({
                'status': 'ERROR',
                'msg': f"Insufficient stock! Available stock for '{item.name}' is only {item.available_stock} {item.unit}."
            })

        if issue_date_str:
            try:
                issue_date = datetime.datetime.strptime(issue_date_str, '%Y-%m-%d').date()
            except:
                issue_date = timezone.now().date()
        else:
            issue_date = timezone.now().date()

        months = item.warranty_months or 6
        warranty_expiry_date = issue_date + datetime.timedelta(days=int(months * 30))

        existing_active = SafetyEquipmentIssue.objects.filter(employee=employee, item=item, status='ACTIVE').first()
        warn_msg = ''
        if existing_active:
            warn_msg = f" (Note: {employee.name} already had an active issue from {existing_active.issue_date})."

        item.available_stock = max(0, item.available_stock - quantity)
        item.save()

        issue = SafetyEquipmentIssue.objects.create(
            employee=employee,
            item=item,
            quantity=quantity,
            size_specification=size_specification,
            issue_date=issue_date,
            warranty_expiry_date=warranty_expiry_date,
            status='ACTIVE',
            issued_by=request.user,
            remarks=remarks
        )

        return JsonResponse({
            'status': 'SUCCESS',
            'msg': f"Successfully issued {quantity}x {item.name} to {employee.name} (Warranty valid till {warranty_expiry_date}){warn_msg}."
        })
    except Exception as e:
        return JsonResponse({'status': 'ERROR', 'msg': str(e)})


@login_required
def api_safety_issue_edit(request, issue_id):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    issue = get_object_or_404(SafetyEquipmentIssue, id=issue_id)
    try:
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST
        size_specification = data.get('size_specification', issue.size_specification)
        status = data.get('status', issue.status)
        remarks = data.get('remarks', issue.remarks)
        issue_date_str = data.get('issue_date')
        expiry_date_str = data.get('warranty_expiry_date')

        old_status = issue.status
        issue.size_specification = size_specification
        issue.remarks = remarks

        if issue_date_str:
            try:
                issue.issue_date = datetime.datetime.strptime(issue_date_str, '%Y-%m-%d').date()
            except:
                pass

        if expiry_date_str:
            try:
                issue.warranty_expiry_date = datetime.datetime.strptime(expiry_date_str, '%Y-%m-%d').date()
            except:
                pass

        if old_status != status:
            issue.status = status
            if old_status in ['VOID', 'RETURNED_RESIGNED'] and status == 'ACTIVE':
                if issue.item.available_stock >= issue.quantity:
                    issue.item.available_stock = max(0, issue.item.available_stock - issue.quantity)
                    issue.item.save()
            elif old_status == 'ACTIVE' and status in ['VOID', 'RETURNED_RESIGNED']:
                issue.item.available_stock += issue.quantity
                issue.item.save()

        issue.save()
        return JsonResponse({'status': 'SUCCESS', 'msg': 'Issue record updated successfully.'})
    except Exception as e:
        return JsonResponse({'status': 'ERROR', 'msg': str(e)})


@login_required
def api_safety_issue_delete(request, issue_id):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    issue = get_object_or_404(SafetyEquipmentIssue, id=issue_id)
    if issue.status == 'ACTIVE':
        issue.item.available_stock += issue.quantity
        issue.item.save()

    emp_name = issue.employee.name if issue.employee else 'Worker'
    item_name = issue.item.name if issue.item else 'Item'
    issue.delete()
    return JsonResponse({'status': 'SUCCESS', 'msg': f"Deleted issue record for {emp_name} - {item_name}. Stock restored."})


def api_safety_employee_search(request):
    """
    Ultra-fast AJAX endpoint to search across all workers by name, emp_id, cid, designation, contractor, or department.
    """
    if not request.user.is_authenticated:
        return JsonResponse({'status': 'ERROR', 'msg': 'Not authenticated.'}, status=401)

    q = request.GET.get('q', '').strip()
    dept = request.GET.get('dept', '').strip()

    qs = Employee.objects.all()
    if dept:
        qs = qs.filter(department__iexact=dept)
    if q:
        qs = qs.filter(
            Q(name__icontains=q) |
            Q(emp_id__icontains=q) |
            Q(cid_number__icontains=q) |
            Q(designation__icontains=q) |
            Q(contractor_agency__icontains=q)
        )

    results = []
    for emp in qs.order_by('name')[:35]:
        results.append({
            'id': emp.id,
            'name': emp.name,
            'emp_id': emp.emp_id or 'N/A',
            'department': emp.department or '',
            'designation': emp.designation or '',
            'contractor_agency': emp.contractor_agency or '',
            'cid_number': emp.cid_number or '',
        })

    return JsonResponse({'status': 'SUCCESS', 'employees': results})


@login_required
def api_safety_worker_active_issues(request, employee_id):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    employee = get_object_or_404(Employee, id=employee_id)
    active_issues = SafetyEquipmentIssue.objects.filter(employee=employee, status='ACTIVE').select_related('item').order_by('-issue_date')

    today = timezone.now().date()
    items_data = []
    for iss in active_issues:
        is_under = iss.warranty_expiry_date and (today <= iss.warranty_expiry_date)
        days_rem = max(0, (iss.warranty_expiry_date - today).days) if iss.warranty_expiry_date else 0
        warranty_months = iss.item.warranty_months or 6
        warranty_days = warranty_months * 30

        # Calculate days used: if issue_date is in the future (bad imported data),
        # derive from expiry date backwards for accurate months_used
        if iss.issue_date and iss.issue_date <= today:
            days_used = (today - iss.issue_date).days
        elif iss.warranty_expiry_date:
            # Use expiry date to back-calculate: days_used = warranty_days - days_remaining
            days_used = max(0, warranty_days - days_rem)
        else:
            days_used = 0

        months_used = max(0, round(days_used / 30))
        remaining_months = max(0, warranty_months - months_used)
        full_cost = float(iss.item.fine_amount or 0.00)
        
        # Prorated fine formula matching company policy: (Full Cost / Life) * Remaining Months
        if remaining_months > 0 and warranty_months > 0:
            prorated_fine = round((full_cost / warranty_months) * remaining_months, 2)
        else:
            prorated_fine = 0.0

        items_data.append({
            'issue_id': iss.id,
            'item_id': iss.item.id,
            'item_name': iss.item.name,
            'item_code': iss.item.item_code,
            'category': iss.item.category,
            'size': iss.size_specification or 'Standard',
            'quantity': iss.quantity,
            'issue_date': iss.issue_date.strftime('%d-%b-%Y') if iss.issue_date else '',
            'issue_date_iso': iss.issue_date.strftime('%Y-%m-%d') if iss.issue_date else '',
            'warranty_expiry_date': iss.warranty_expiry_date.strftime('%d-%b-%Y') if iss.warranty_expiry_date else '',
            'warranty_expiry_date_iso': iss.warranty_expiry_date.strftime('%Y-%m-%d') if iss.warranty_expiry_date else '',
            'is_under_warranty': is_under,
            'days_remaining': days_rem,
            'days_used': days_used,
            'months_used': months_used,
            'remaining_months': remaining_months,
            'warranty_months': warranty_months,
            'full_cost': full_cost,
            'prorated_fine': prorated_fine,
            'standard_fine': full_cost,
            'available_stock': iss.item.available_stock,
            'unit': iss.item.unit,
            'issued_by': (iss.issued_by.get_full_name() or iss.issued_by.username) if iss.issued_by else 'Admin'
        })

    return JsonResponse({
        'status': 'SUCCESS',
        'employee_name': employee.name,
        'emp_id': employee.emp_id,
        'department': employee.department,
        'active_issues': items_data
    })


# ==============================================================================
# SMART RETURN & REPLACEMENT DESK WITH WARRANTY & FINE ENFORCEMENT
# ==============================================================================

@login_required
def api_safety_process_replacement(request):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    try:
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST
        issue_id = data.get('original_issue_id')
        reason = data.get('reason', 'DAMAGED_PREMATURE')
        old_condition = data.get('old_item_condition', '').strip()
        issue_new_item = str(data.get('issue_new_item', 'true')).lower() in ['true', '1', 'yes']
        new_size = data.get('new_size', '').strip()
        fine_status = data.get('fine_status', 'APPROVED_SALARY_DEDUCTION')
        fine_override = data.get('fine_override')
        waived_reason = data.get('waived_reason', '').strip()

        original_issue = get_object_or_404(SafetyEquipmentIssue, id=issue_id)
        employee = original_issue.employee
        item = original_issue.item

        today = timezone.now().date()
        is_under_warranty = original_issue.warranty_expiry_date and (today <= original_issue.warranty_expiry_date)
        warranty_months = item.warranty_months or 6
        warranty_days = warranty_months * 30
        days_rem = max(0, (original_issue.warranty_expiry_date - today).days) if original_issue.warranty_expiry_date else 0

        # Handle bad imported data where issue_date may be in the future
        if original_issue.issue_date and original_issue.issue_date <= today:
            days_used = (today - original_issue.issue_date).days
        elif original_issue.warranty_expiry_date:
            days_used = max(0, warranty_days - days_rem)
        else:
            days_used = 0

        months_used = max(0, round(days_used / 30))
        remaining_months = max(0, warranty_months - months_used)
        full_cost = item.fine_amount or Decimal('0.00')

        if reason in ['WORN_OUT_NORMAL', 'SIZE_EXCHANGE'] or remaining_months == 0 or not is_under_warranty:
            is_premature = False
            fine_amount = Decimal('0.00')
            fine_status = 'NO_FINE'
            new_issue_status = 'REPLACED_NORMAL'
        else:
            is_premature = True
            new_issue_status = 'REPLACED_PREMATURE'

            if fine_override is not None and str(fine_override).strip() != '':
                try:
                    fine_amount = max(Decimal('0.00'), Decimal(str(fine_override)))
                except:
                    fine_amount = round(Decimal(str(full_cost)) / Decimal(str(warranty_months)) * Decimal(str(remaining_months)), 2)
            else:
                # Prorated fine formula matching company policy:
                # (Full Cost of PPE / Life in Months) * Remaining Months
                if remaining_months > 0 and warranty_months > 0:
                    fine_amount = round(Decimal(str(full_cost)) / Decimal(str(warranty_months)) * Decimal(str(remaining_months)), 2)
                else:
                    fine_amount = Decimal('0.00')

            if fine_status == 'WAIVED':
                fine_amount = Decimal('0.00')
            elif not fine_status or fine_status == 'PENDING_APPROVAL':
                fine_status = 'APPROVED_SALARY_DEDUCTION'

        original_issue.status = new_issue_status
        original_issue.save()

        new_issue = None
        if issue_new_item:
            if item.available_stock < 1:
                return JsonResponse({
                    'status': 'ERROR',
                    'msg': f"Cannot issue fresh replacement: '{item.name}' is out of stock! Old return logged, please restock item first."
                })

            item.available_stock = max(0, item.available_stock - 1)
            item.save()

            fresh_warranty_expiry = today + datetime.timedelta(days=int(warranty_months * 30))
            new_issue = SafetyEquipmentIssue.objects.create(
                employee=employee,
                item=item,
                quantity=1,
                size_specification=new_size or original_issue.size_specification,
                issue_date=today,
                warranty_expiry_date=fresh_warranty_expiry,
                status='ACTIVE',
                issued_by=request.user,
                remarks=f"Replaced from Issue #{original_issue.id} ({reason})"
            )

        replacement_log = SafetyReplacementAndFine.objects.create(
            original_issue=original_issue,
            new_issue=new_issue,
            employee=employee,
            item=item,
            replacement_date=today,
            reason=reason,
            old_item_condition=old_condition,
            is_premature=is_premature,
            days_used=days_used,
            warranty_days=warranty_days,
            months_used=months_used,
            remaining_months=remaining_months,
            full_cost=full_cost,
            fine_amount=fine_amount,
            fine_status=fine_status,
            waived_reason=waived_reason if fine_status == 'WAIVED' else None,
            processed_by=request.user
        )

        msg_fine = f"Fine logged: Nu. {fine_amount} ({fine_status})" if fine_amount > 0 else "Replacement issued Free of Cost (No Fine)."
        return JsonResponse({
            'status': 'SUCCESS',
            'msg': f"Replacement processed for {employee.name}. {msg_fine}"
        })
    except Exception as e:
        return JsonResponse({'status': 'ERROR', 'msg': str(e)})


@login_required
def api_safety_fine_update_status(request, fine_id):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    fine = get_object_or_404(SafetyReplacementAndFine, id=fine_id)
    try:
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST
        new_status = data.get('fine_status', fine.fine_status)
        waived_reason = data.get('waived_reason', fine.waived_reason)
        fine_override = data.get('fine_amount')

        fine.fine_status = new_status
        if waived_reason:
            fine.waived_reason = waived_reason
        if fine_override is not None and str(fine_override).strip() != '':
            try:
                fine.fine_amount = max(Decimal('0.00'), Decimal(str(fine_override)))
            except:
                pass

        if new_status == 'WAIVED' and not fine.fine_amount:
            fine.fine_amount = Decimal('0.00')

        fine.save()
        return JsonResponse({'status': 'SUCCESS', 'msg': f"Fine status updated to {fine.get_fine_status_display()}."})
    except Exception as e:
        return JsonResponse({'status': 'ERROR', 'msg': str(e)})


@login_required
def api_safety_fine_delete(request, fine_id):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    fine = get_object_or_404(SafetyReplacementAndFine, id=fine_id)
    emp_name = fine.employee.name if fine.employee else 'Worker'
    fine.delete()
    return JsonResponse({'status': 'SUCCESS', 'msg': f"Fine & replacement record for {emp_name} deleted."})


# ==============================================================================
# EXPORT TO EXCEL: FINES REGISTER, INVENTORY, ISSUANCES
# ==============================================================================

@login_required
def export_safety_fines_excel(request):
    if not has_safety_access(request.user):
        return HttpResponse('Permission denied', status=403)

    today = timezone.now().date()
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    dept = request.GET.get('department')
    reason = request.GET.get('reason')
    status = request.GET.get('status')

    fines = SafetyReplacementAndFine.objects.select_related('employee', 'item', 'original_issue', 'processed_by').all()

    if from_date:
        fines = fines.filter(replacement_date__gte=from_date)
    if to_date:
        fines = fines.filter(replacement_date__lte=to_date)
    if dept and dept != 'ALL':
        fines = fines.filter(employee__department__icontains=dept)
    if reason and reason != 'ALL':
        fines = fines.filter(reason=reason)
    if status and status != 'ALL':
        fines = fines.filter(fine_status=status)

    fines = fines.order_by('-replacement_date')

    # Build exact format matching Safety_Materials_Deduction_Report_Format.docx
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Deduction Report'
    ws.views.sheetView[0].showGridLines = True

    report_no = f"SMDR-{today.strftime('%Y%m%d')}-{len(fines):03d}"
    report_date = today.strftime('%d-%m-%Y')
    period_str = f"{from_date} to {to_date}" if (from_date and to_date) else f"01-01-{today.year} to {today.strftime('%d-%m-%Y')}"
    dept_site = dept if (dept and dept != 'ALL') else "Safety Department / Site Operations"
    prep_by = request.user.get_full_name() or request.user.username or "Admin"
    desig = "Safety Officer / HSE Manager"

    # Styles
    font_title = Font(name='Calibri', size=15, bold=True, color='FFFFFF')
    fill_title = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')

    font_meta_lbl = Font(name='Calibri', size=10, bold=True, color='1E293B')
    font_meta_val = Font(name='Calibri', size=10, bold=False, color='0F172A')

    font_th = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    fill_th = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')

    font_data = Font(name='Calibri', size=9.5, color='0F172A')
    font_data_bold = Font(name='Calibri', size=9.5, bold=True, color='0F172A')
    font_amt = Font(name='Calibri', size=10, bold=True, color='B91C1C')
    font_tot = Font(name='Calibri', size=11, bold=True, color='0F172A')

    fill_alt = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_tot = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

    thin_border_side = Side(border_style='thin', color='CBD5E1')
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    double_bottom_side = Side(border_style='double', color='0F172A')
    top_thick_side = Side(border_style='medium', color='0F172A')
    tot_border = Border(left=thin_border_side, right=thin_border_side, top=top_thick_side, bottom=double_bottom_side)

    # 1. Title Banner (Rows 1-2 merged A1:O2)
    ws.merge_cells('A1:O2')
    title_cell = ws['A1']
    title_cell.value = "SAFETY MATERIALS DEDUCTION REPORT"
    title_cell.font = font_title
    title_cell.fill = fill_title
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    # 2. Metadata Block (Rows 4-5)
    ws['B4'] = "Report No.:"
    ws['B4'].font = font_meta_lbl
    ws['C4'] = report_no
    ws['C4'].font = font_meta_val

    ws['F4'] = "Report Date:"
    ws['F4'].font = font_meta_lbl
    ws['G4'] = report_date
    ws['G4'].font = font_meta_val

    ws['K4'] = "Department / Site:"
    ws['K4'].font = font_meta_lbl
    ws['L4'] = dept_site
    ws['L4'].font = font_meta_val

    ws['B5'] = "Period Covered:"
    ws['B5'].font = font_meta_lbl
    ws['C5'] = period_str
    ws['C5'].font = font_meta_val

    ws['F5'] = "Prepared By:"
    ws['F5'].font = font_meta_lbl
    ws['G5'] = prep_by
    ws['G5'].font = font_meta_val

    ws['K5'] = "Designation:"
    ws['K5'].font = font_meta_lbl
    ws['L5'] = desig
    ws['L5'].font = font_meta_val

    # 3. Table Headers (Row 7)
    headers = [
        'Sl. No.',
        'Employee ID',
        'Employee Name',
        'Designation / Section',
        'Contractor / Hiring Agency',
        'Material Issued',
        'Original Issue Date',
        'Standard Lifeline (Months)',
        'Due Expiry Date',
        'Re-issue Date',
        'Period Used (Months)',
        'Unused Period (Months)',
        'Unit Cost (Nu.)',
        'Deduction Amount (Nu.)',
        'Remarks'
    ]

    header_row = 7
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.value = h_text
        cell.font = font_th
        cell.fill = fill_th
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = data_border
    ws.row_dimensions[header_row].height = 28

    # 4. Data Rows
    cur_row = 8
    total_deduction = 0.0

    for idx, f in enumerate(fines, start=1):
        emp = f.employee
        item = f.item
        orig_iss = f.original_issue

        emp_id = emp.emp_id if emp else 'N/A'
        emp_name = emp.name if emp else 'N/A'
        emp_desig = (emp.designation or emp.department) if emp else 'Site Worker'
        emp_agency = (emp.contractor_agency or emp.department) if emp else 'General'
        mat_name = item.name if item else 'Equipment'

        lifeline_m = item.warranty_months if (item and item.warranty_months) else (round(f.warranty_days / 30) or 6)

        if orig_iss and orig_iss.issue_date:
            orig_date_str = orig_iss.issue_date.strftime('%d-%b-%Y')
            orig_date_raw = orig_iss.issue_date
        else:
            orig_date_raw = f.replacement_date - datetime.timedelta(days=max(0, f.months_used * 30))
            orig_date_str = orig_date_raw.strftime('%d-%b-%Y')

        if orig_iss and orig_iss.warranty_expiry_date:
            due_date_str = orig_iss.warranty_expiry_date.strftime('%d-%b-%Y')
        else:
            due_date_raw = orig_date_raw + datetime.timedelta(days=lifeline_m * 30)
            due_date_str = due_date_raw.strftime('%d-%b-%Y')

        reissue_str = f.replacement_date.strftime('%d-%b-%Y') if f.replacement_date else today.strftime('%d-%b-%Y')

        used_m = f.months_used
        unused_m = f.remaining_months if f.remaining_months is not None else max(0, lifeline_m - used_m)
        unit_cost = float(f.full_cost or (item.fine_amount if item else 0.0) or 0.0)
        ded_amt = float(f.fine_amount or 0.0)
        total_deduction += ded_amt

        remarks_list = []
        if f.reason:
            remarks_list.append(f.get_reason_display())
        if f.old_item_condition:
            remarks_list.append(f"Cond: {f.old_item_condition}")
        if f.waived_reason:
            remarks_list.append(f"Waived: {f.waived_reason}")
        remarks_str = ' | '.join(remarks_list) or 'Premature Replacement'

        row_data = [
            idx,
            emp_id,
            emp_name,
            emp_desig,
            emp_agency,
            mat_name,
            orig_date_str,
            lifeline_m,
            due_date_str,
            reissue_str,
            used_m,
            unused_m,
            round(unit_cost, 2),
            round(ded_amt, 2),
            remarks_str
        ]

        is_even = (idx % 2 == 0)
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.value = val
            cell.font = font_amt if col_idx == 14 and ded_amt > 0 else (font_data_bold if col_idx in [1, 2, 3] else font_data)
            if is_even:
                cell.fill = fill_alt
            cell.border = data_border

            if col_idx in [1, 7, 8, 9, 10, 11, 12]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in [13, 14]:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0.00'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

        ws.row_dimensions[cur_row].height = 20
        cur_row += 1

    # 5. Total Row
    tot_row = cur_row
    ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=13)
    tot_lbl = ws.cell(row=tot_row, column=1)
    tot_lbl.value = "TOTAL DEDUCTION AMOUNT (Nu.)"
    tot_lbl.font = font_tot
    tot_lbl.alignment = Alignment(horizontal='right', vertical='center')

    for col in range(1, 14):
        c = ws.cell(row=tot_row, column=col)
        c.fill = fill_tot
        c.border = tot_border

    tot_val = ws.cell(row=tot_row, column=14)
    tot_val.value = total_deduction
    tot_val.font = font_amt
    tot_val.number_format = '#,##0.00'
    tot_val.fill = fill_tot
    tot_val.border = tot_border
    tot_val.alignment = Alignment(horizontal='right', vertical='center')

    tot_end = ws.cell(row=tot_row, column=15)
    tot_end.value = ""
    tot_end.fill = fill_tot
    tot_end.border = tot_border

    ws.row_dimensions[tot_row].height = 24

    # 6. Signature Block
    sig_row_title = tot_row + 3
    sig_row_line = sig_row_title + 1

    ws.cell(row=sig_row_title, column=2).value = "Prepared By:"
    ws.cell(row=sig_row_title, column=2).font = font_meta_lbl
    ws.cell(row=sig_row_line, column=2).value = "Name / Signature / Date"
    ws.cell(row=sig_row_line, column=2).font = Font(name='Calibri', size=9, italic=True, color='64748B')

    ws.cell(row=sig_row_title, column=7).value = "Checked By:"
    ws.cell(row=sig_row_title, column=7).font = font_meta_lbl
    ws.cell(row=sig_row_line, column=7).value = "Name / Signature / Date"
    ws.cell(row=sig_row_line, column=7).font = Font(name='Calibri', size=9, italic=True, color='64748B')

    ws.cell(row=sig_row_title, column=12).value = "Approved By:"
    ws.cell(row=sig_row_title, column=12).font = font_meta_lbl
    ws.cell(row=sig_row_line, column=12).value = "Name / Signature / Date"
    ws.cell(row=sig_row_line, column=12).font = Font(name='Calibri', size=9, italic=True, color='64748B')

    col_widths = {
        1: 8, 2: 15, 3: 24, 4: 26, 5: 28, 6: 22, 7: 15, 8: 14,
        9: 15, 10: 15, 11: 14, 12: 14, 13: 15, 14: 16, 15: 28
    }
    for c_idx, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(c_idx)].width = width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=Safety_Materials_Deduction_Report_{today.strftime("%Y%m%d")}.xlsx'
    wb.save(response)
    return response

def export_safety_inventory_excel(request):
    if not has_safety_access(request.user):
        return HttpResponse('Permission denied', status=403)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Safety Store Inventory'

    headers = [
        'SL', 'ITEM CODE', 'EQUIPMENT NAME', 'CATEGORY', 'UNIT', 'SPECIFICATIONS',
        'TOTAL STOCK', 'AVAILABLE STOCK', 'ISSUED TO WORKERS', 'WARRANTY (MONTHS)',
        'STANDARD COST / FINE RATE (Nu)', 'MIN ALERT LEVEL', 'LOW STOCK STATUS', 'NOTES'
    ]

    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    items = SafetyStoreItem.objects.all()

    category = request.GET.get('category')
    stock_filter = request.GET.get('stock')

    if category and category != 'ALL':
        items = items.filter(category=category)
    if stock_filter == 'LOW':
        items = [i for i in items if i.is_low_stock]
    elif stock_filter == 'IN_STOCK':
        items = items.filter(available_stock__gt=0)

    if not isinstance(items, list):
        items = items.order_by('category', 'name')
    else:
        items = sorted(items, key=lambda x: (x.category, x.name))

    for idx, item in enumerate(items, 1):
        row = [
            idx,
            item.item_code,
            item.name,
            item.category,
            item.unit,
            item.specification or '',
            item.total_stock,
            item.available_stock,
            item.issued_stock,
            item.warranty_months,
            float(item.fine_amount),
            item.minimum_alert_level,
            'LOW STOCK' if item.is_low_stock else 'NORMAL',
            item.notes or ''
        ]
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=Safety_Store_Inventory_{timezone.now().strftime("%Y%m%d")}.xlsx'
    wb.save(response)
    return response


@login_required
def export_safety_issues_excel(request):
    if not has_safety_access(request.user):
        return HttpResponse('Permission denied', status=403)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Safety PPE Issuances'

    headers = [
        'SL', 'EMPLOYEE ID', 'EMPLOYEE NAME', 'DEPARTMENT', 'EQUIPMENT NAME', 'ITEM CODE',
        'QTY', 'SIZE / SPEC', 'ISSUE DATE', 'WARRANTY EXPIRY', 'WARRANTY STATUS', 'STATUS', 'ISSUED BY', 'REMARKS'
    ]

    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    issues = SafetyEquipmentIssue.objects.select_related('employee', 'item', 'issued_by').all()

    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    dept = request.GET.get('department')
    category = request.GET.get('category')
    status = request.GET.get('status')

    if from_date:
        issues = issues.filter(issue_date__gte=from_date)
    if to_date:
        issues = issues.filter(issue_date__lte=to_date)
    if dept and dept != 'ALL':
        issues = issues.filter(employee__department__icontains=dept)
    if category and category != 'ALL':
        issues = issues.filter(item__category=category)
    if status and status != 'ALL':
        issues = issues.filter(status=status)

    issues = issues.order_by('-issue_date')
    today = timezone.now().date()

    for idx, iss in enumerate(issues, 1):
        is_under = iss.warranty_expiry_date and (today <= iss.warranty_expiry_date)
        w_status = f'Active ({iss.days_remaining}d left)' if is_under else 'Expired'
        row = [
            idx,
            iss.employee.emp_id if iss.employee else '',
            iss.employee.name if iss.employee else '',
            iss.employee.department if iss.employee else '',
            iss.item.name if iss.item else '',
            iss.item.item_code if iss.item else '',
            iss.quantity,
            iss.size_specification or '',
            iss.issue_date.strftime('%Y-%m-%d') if iss.issue_date else '',
            iss.warranty_expiry_date.strftime('%Y-%m-%d') if iss.warranty_expiry_date else '',
            w_status,
            iss.get_status_display(),
            iss.issued_by.get_full_name() or iss.issued_by.username if iss.issued_by else '',
            iss.remarks or ''
        ]
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=Safety_Equipment_Issuances_{timezone.now().strftime("%Y%m%d")}.xlsx'
    wb.save(response)
    return response


# ==============================================================================
# UNIFIED NATIVE PDF EXPORT FOR SAFETY DEPARTMENT
# ==============================================================================

@login_required
def export_safety_report_pdf(request):
    if not has_safety_access(request.user):
        return HttpResponse('Permission denied', status=403)

    report_type = request.GET.get('report_type', 'inventory')
    format_type = request.GET.get('format', 'print')  # 'print' or 'pdf'
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    dept = request.GET.get('department')
    category = request.GET.get('category')
    status = request.GET.get('status')

    today_str = timezone.now().strftime('%d %B %Y, %I:%M %p')

    # Prepare title & data based on report_type
    if report_type == 'inventory':
        report_title = "Safety Department - Store Inventory & Stock Catalog"
        items = SafetyStoreItem.objects.all()
        if category and category != 'ALL': items = items.filter(category=category)
        if status == 'LOW': items = [i for i in items if i.is_low_stock]
        elif status == 'IN_STOCK': items = items.filter(available_stock__gt=0)
        if not isinstance(items, list): items = items.order_by('category', 'name')
        else: items = sorted(items, key=lambda x: (x.category, x.name))

        total_avail = sum(i.available_stock for i in items)
        total_stk = sum(i.total_stock for i in items)

        rows_html = ""
        for idx, item in enumerate(items, 1):
            st_color = "#dc2626" if item.is_low_stock else "#16a34a"
            st_label = "Low Stock" if item.is_low_stock else "Available"
            rows_html += f'''
            <tr>
                <td style="text-align:center;">{idx}</td>
                <td style="font-family:monospace;font-weight:bold;color:#2563eb;">{item.item_code}</td>
                <td><b>{item.name}</b><br><small style="color:#64748b;">{item.specification or ''}</small></td>
                <td><span class="badge">{item.category}</span></td>
                <td style="text-align:center;"><b>{item.available_stock}</b> / {item.total_stock} {item.unit}</td>
                <td style="text-align:center;">{item.warranty_months} Mos</td>
                <td style="text-align:right;">Nu. {float(item.fine_amount):.2f}</td>
                <td style="text-align:center;color:{st_color};font-weight:bold;">{st_label}</td>
            </tr>'''

        table_headers = "<th>#</th><th>Item Code</th><th>Equipment Name</th><th>Category</th><th>Stock (Avail/Total)</th><th>Warranty</th><th>Standard Rate</th><th>Status</th>"
        summary_kpi = f"<b>Total Items:</b> {len(items)} &bull; <b>Available Quantity:</b> {total_avail} &bull; <b>Total Stock:</b> {total_stk}"

    elif report_type == 'issues':
        report_title = "Safety Department - Equipment Allocation & Issuance Register"
        issues = SafetyEquipmentIssue.objects.select_related('employee', 'item').all()
        if from_date: issues = issues.filter(issue_date__gte=from_date)
        if to_date: issues = issues.filter(issue_date__lte=to_date)
        if dept and dept != 'ALL': issues = issues.filter(employee__department__icontains=dept)
        if category and category != 'ALL': issues = issues.filter(item__category=category)
        if status and status != 'ALL': issues = issues.filter(status=status)
        issues = issues.order_by('-issue_date')[:2000]

        rows_html = ""
        today = timezone.now().date()
        for idx, iss in enumerate(issues, 1):
            is_under = iss.warranty_expiry_date and (today <= iss.warranty_expiry_date)
            w_color = "#16a34a" if is_under else "#dc2626"
            w_text = f"Active ({iss.days_remaining}d left)" if is_under else "Expired"
            rows_html += f'''
            <tr>
                <td style="text-align:center;">{idx}</td>
                <td><b>{iss.employee.name if iss.employee else 'N/A'}</b><br><small style="font-family:monospace;color:#2563eb;">{iss.employee.emp_id if iss.employee else ''}</small></td>
                <td>{iss.employee.department if iss.employee else '-'}</td>
                <td><b>{iss.item.name if iss.item else 'N/A'}</b><br><small style="color:#64748b;">({iss.size_specification or 'Standard'})</small></td>
                <td style="text-align:center;">{iss.issue_date.strftime('%d-%b-%Y') if iss.issue_date else '-'}</td>
                <td style="text-align:center;">{iss.warranty_expiry_date.strftime('%d-%b-%Y') if iss.warranty_expiry_date else '-'}</td>
                <td style="text-align:center;color:{w_color};font-weight:bold;">{w_text}</td>
                <td style="text-align:center;">{iss.get_status_display()}</td>
            </tr>'''

        table_headers = "<th>#</th><th>Worker Name & ID</th><th>Department</th><th>Equipment / Spec</th><th>Issue Date</th><th>Expiry Date</th><th>Warranty</th><th>Status</th>"
        summary_kpi = f"<b>Total Records:</b> {len(issues)} Issuances"

    else: # fines - formatted per Safety_Materials_Deduction_Report_Format.docx
        today_date = timezone.now().date()
        fines = SafetyReplacementAndFine.objects.select_related('employee', 'item', 'original_issue', 'processed_by').all()
        if from_date: fines = fines.filter(replacement_date__gte=from_date)
        if to_date: fines = fines.filter(replacement_date__lte=to_date)
        if dept and dept != 'ALL': fines = fines.filter(employee__department__icontains=dept)
        if status and status != 'ALL': fines = fines.filter(fine_status=status)
        fines = fines.order_by('-replacement_date')[:2000]

        total_deduction = sum(f.fine_amount for f in fines)
        report_no = f"SMDR-{today_date.strftime('%Y%m%d')}-{len(fines):03d}"
        period_str = f"{from_date} to {to_date}" if (from_date and to_date) else f"01-Jan-{today_date.year} to {today_date.strftime('%d-%b-%Y')}"
        dept_str = dept if (dept and dept != 'ALL') else "Safety Department / Site Operations"
        prep_by = request.user.get_full_name() or request.user.username or "Admin"
        desig_str = "Safety Officer / HSE Manager"

        rows_html = ""
        for idx, f in enumerate(fines, 1):
            emp = f.employee
            item = f.item
            orig_iss = f.original_issue

            emp_id = emp.emp_id if emp else 'N/A'
            emp_name = emp.name if emp else 'N/A'
            emp_desig = (emp.designation or emp.department) if emp else 'Site Worker'
            emp_agency = (emp.contractor_agency or emp.department) if emp else 'General'
            mat_name = item.name if item else 'Equipment'
            lifeline_m = item.warranty_months if (item and item.warranty_months) else (round(f.warranty_days / 30) or 6)

            if orig_iss and orig_iss.issue_date:
                orig_date_str = orig_iss.issue_date.strftime('%d-%b-%Y')
                orig_date_raw = orig_iss.issue_date
            else:
                orig_date_raw = f.replacement_date - datetime.timedelta(days=max(0, f.months_used * 30))
                orig_date_str = orig_date_raw.strftime('%d-%b-%Y')

            if orig_iss and orig_iss.warranty_expiry_date:
                due_date_str = orig_iss.warranty_expiry_date.strftime('%d-%b-%Y')
            else:
                due_date_raw = orig_date_raw + datetime.timedelta(days=lifeline_m * 30)
                due_date_str = due_date_raw.strftime('%d-%b-%Y')

            reissue_str = f.replacement_date.strftime('%d-%b-%Y') if f.replacement_date else today_date.strftime('%d-%b-%Y')
            used_m = f.months_used
            unused_m = f.remaining_months if f.remaining_months is not None else max(0, lifeline_m - used_m)
            unit_cost = float(f.full_cost or (item.fine_amount if item else 0.0) or 0.0)
            ded_amt = float(f.fine_amount or 0.0)

            remarks_list = []
            if f.reason: remarks_list.append(f.get_reason_display())
            if f.old_item_condition: remarks_list.append(f"Cond: {f.old_item_condition}")
            if f.waived_reason: remarks_list.append(f"Waived: {f.waived_reason}")
            remarks_str = ' | '.join(remarks_list) or 'Premature Replacement'

            rows_html += f'''
            <tr>
                <td style="text-align:center;">{idx}</td>
                <td style="font-family:monospace;font-weight:700;color:#1e40af;">{emp_id}</td>
                <td><b>{emp_name}</b></td>
                <td>{emp_desig}</td>
                <td>{emp_agency}</td>
                <td><b>{mat_name}</b></td>
                <td style="text-align:center;">{orig_date_str}</td>
                <td style="text-align:center;">{lifeline_m}</td>
                <td style="text-align:center;">{due_date_str}</td>
                <td style="text-align:center;">{reissue_str}</td>
                <td style="text-align:center;">{used_m}</td>
                <td style="text-align:center;">{unused_m}</td>
                <td style="text-align:right;">{unit_cost:.2f}</td>
                <td style="text-align:right;font-weight:800;color:#b91c1c;">{ded_amt:.2f}</td>
                <td style="font-size:8.5px;color:#475569;">{remarks_str}</td>
            </tr>'''

        fines_html = f'''<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>SAFETY MATERIALS DEDUCTION REPORT</title>
    <style>
        @page {{ size: A4 landscape; margin: 8mm 10mm 10mm 10mm; }}
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif; font-size: 10px; color: #0f172a; margin: 0; padding: 8px; background: #ffffff; }}
        .report-title-banner {{ background: #0f172a; color: #ffffff; text-align: center; padding: 10px 16px; border-radius: 6px; margin-bottom: 12px; }}
        .report-title-banner h1 {{ margin: 0; font-size: 16px; font-weight: 800; letter-spacing: 0.5px; }}
        .meta-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 8px; background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 6px; padding: 10px 14px; margin-bottom: 12px; font-size: 10px; }}
        .meta-row {{ display: flex; margin-bottom: 4px; }}
        .meta-lbl {{ font-weight: 700; color: #475569; width: 140px; }}
        .meta-val {{ font-weight: 600; color: #0f172a; }}
        table {{ width: 100%; border-collapse: collapse; font-size: 9px; margin-bottom: 12px; }}
        th {{ background: #1e3a8a; color: #ffffff; font-weight: 700; padding: 6px 4px; border: 1px solid #1e3a8a; text-align: center; vertical-align: middle; }}
        td {{ padding: 5px 4px; border: 1px solid #cbd5e1; vertical-align: middle; }}
        tr:nth-child(even) {{ background: #f8fafc; }}
        .total-row {{ background: #e2e8f0 !important; font-weight: 800; font-size: 10px; }}
        .total-row td {{ border-top: 2px solid #0f172a; border-bottom: 3px double #0f172a; }}
        .formula-box {{ background: #fefce8; border: 1px solid #fde047; border-radius: 4px; padding: 8px 12px; margin-bottom: 20px; font-size: 9.5px; color: #713f12; line-height: 1.4; }}
        .sig-grid {{ display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 24px; margin-top: 24px; page-break-inside: avoid; }}
        .sig-box {{ border-top: 1.5px solid #475569; padding-top: 6px; text-align: center; }}
        .sig-title {{ font-weight: 800; font-size: 10px; color: #0f172a; margin-bottom: 2px; }}
        .sig-sub {{ font-size: 8.5px; color: #64748b; font-style: italic; }}
        .print-btn-bar {{ margin-bottom: 12px; display: flex; justify-content: flex-end; gap: 8px; }}
        @media print {{ .print-btn-bar {{ display: none !important; }} }}
    </style>
</head>
<body>
    <div class="print-btn-bar">
        <button onclick="window.print()" style="background:#2563eb;color:#ffffff;border:none;padding:7px 16px;border-radius:6px;font-weight:bold;cursor:pointer;">
            🖨️ Print / Save as PDF
        </button>
    </div>

    <div class="report-title-banner">
        <h1>SAFETY MATERIALS DEDUCTION REPORT</h1>
    </div>

    <div class="meta-grid">
        <div>
            <div class="meta-row"><span class="meta-lbl">Report No.:</span><span class="meta-val">{report_no}</span></div>
            <div class="meta-row"><span class="meta-lbl">Period Covered:</span><span class="meta-val">{period_str}</span></div>
            <div class="meta-row"><span class="meta-lbl">Department / Site:</span><span class="meta-val">{dept_str}</span></div>
        </div>
        <div>
            <div class="meta-row"><span class="meta-lbl">Report Date:</span><span class="meta-val">{today_date.strftime('%d-%m-%Y')}</span></div>
            <div class="meta-row"><span class="meta-lbl">Prepared By:</span><span class="meta-val">{prep_by}</span></div>
            <div class="meta-row"><span class="meta-lbl">Designation:</span><span class="meta-val">{desig_str}</span></div>
        </div>
    </div>

    <table>
        <thead>
            <tr>
                <th style="width:30px;">Sl. No.</th>
                <th style="width:75px;">Employee ID</th>
                <th style="width:110px;">Employee Name</th>
                <th style="width:110px;">Designation / Section</th>
                <th style="width:110px;">Contractor / Hiring Agency</th>
                <th style="width:95px;">Material Issued</th>
                <th style="width:65px;">Original Issue Date</th>
                <th style="width:50px;">Standard Lifeline (Mos)</th>
                <th style="width:65px;">Due Expiry Date</th>
                <th style="width:65px;">Re-issue Date</th>
                <th style="width:45px;">Period Used (Mos)</th>
                <th style="width:45px;">Unused Period (Mos)</th>
                <th style="width:55px;">Unit Cost (Nu.)</th>
                <th style="width:65px;">Deduction Amount (Nu.)</th>
                <th>Remarks</th>
            </tr>
        </thead>
        <tbody>
            {rows_html}
            <tr class="total-row">
                <td colspan="13" style="text-align:right;padding-right:12px;">TOTAL DEDUCTION AMOUNT (Nu.)</td>
                <td style="text-align:right;color:#b91c1c;font-weight:800;">{total_deduction:,.2f}</td>
                <td></td>
            </tr>
        </tbody>
    </table>

    <div class="sig-grid">
        <div class="sig-box">
            <div class="sig-title">Prepared By</div>
            <div class="sig-sub">Name / Signature / Date</div>
        </div>
        <div class="sig-box">
            <div class="sig-title">Checked By</div>
            <div class="sig-sub">Name / Signature / Date</div>
        </div>
        <div class="sig-box">
            <div class="sig-title">Approved By</div>
            <div class="sig-sub">Name / Signature / Date</div>
        </div>
    </div>
</body>
</html>'''
        return HttpResponse(fines_html, content_type='text/html; charset=utf-8')

    html_content = f'''<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{report_title}</title>
    <style>
        @page {{
            size: A4 landscape;
            margin: 10mm 12mm 12mm 12mm;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
            font-size: 11px;
            color: #1e293b;
            margin: 0;
            padding: 10px;
            background: #ffffff;
        }}
        .report-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            border-bottom: 2px solid #0f172a;
            padding-bottom: 10px;
            margin-bottom: 12px;
        }}
        .header-title h1 {{
            margin: 0 0 4px 0;
            font-size: 16px;
            font-weight: 800;
            color: #0f172a;
        }}
        .header-title p {{
            margin: 0;
            font-size: 11px;
            color: #64748b;
        }}
        .header-meta {{
            text-align: right;
            font-size: 10px;
            color: #475569;
        }}
        .kpi-bar {{
            background: #f1f5f9;
            border: 1px solid #cbd5e1;
            border-radius: 6px;
            padding: 8px 12px;
            font-size: 11px;
            margin-bottom: 12px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 10px;
        }}
        th {{
            background: #1e293b;
            color: #ffffff;
            font-weight: 700;
            padding: 7px 6px;
            text-align: left;
            border: 1px solid #334155;
        }}
        td {{
            padding: 6px;
            border: 1px solid #cbd5e1;
            vertical-align: middle;
        }}
        tr:nth-child(even) td {{
            background: #f8fafc;
        }}
        .badge {{
            background: #e2e8f0;
            padding: 2px 5px;
            border-radius: 4px;
            font-weight: 600;
            font-size: 9px;
        }}
        .no-print-bar {{
            background: #2563eb;
            color: #ffffff;
            padding: 10px 16px;
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 14px;
            border-radius: 8px;
        }}
        .btn-print {{
            background: #ffffff;
            color: #1e40af;
            font-weight: 700;
            border: none;
            padding: 6px 14px;
            border-radius: 6px;
            cursor: pointer;
        }}
        @media print {{
            .no-print-bar {{ display: none !important; }}
            body {{ padding: 0; }}
        }}
    </style>
</head>
<body>
    <div class="no-print-bar">
        <div><b>Official Safety Department PDF Report Preview</b> &bull; Ready for download or printing</div>
        <div>
            <button class="btn-print" onclick="window.print()"><i class="fas fa-print"></i> 🖨️ Print / Save as PDF</button>
        </div>
    </div>

    <div class="report-header">
        <div class="header-title">
            <h1>🦺 {report_title}</h1>
            <p>HSE Store Inventory &bull; Worker Equipment Allocation &bull; Warranty &amp; Fine Register</p>
        </div>
        <div class="header-meta">
            <div><b>Generated:</b> {today_str}</div>
            <div><b>Generated By:</b> {request.user.get_full_name() or request.user.username}</div>
        </div>
    </div>

    <div class="kpi-bar">
        {summary_kpi}
        {f' &bull; <b>Department:</b> {dept}' if dept and dept != 'ALL' else ''}
        {f' &bull; <b>Date Range:</b> {from_date} to {to_date}' if from_date or to_date else ''}
    </div>

    <table>
        <thead>
            <tr>{table_headers}</tr>
        </thead>
        <tbody>
            {rows_html if rows_html else '<tr><td colspan="8" style="text-align:center;padding:20px;">No matching records found.</td></tr>'}
        </tbody>
    </table>

    <div style="margin-top: 16px; text-align: right; font-size: 9px; color: #94a3b8;">
        Safety Department Management System &bull; Punatsangchhu Hydroelectric Project &bull; Confidential
    </div>

    <script>
        // Auto prompt print if opened in standalone window
        if (window.location.search.includes('autoprint=1')) {{
            setTimeout(() => window.print(), 500);
        }}
    </script>
</body>
</html>'''

    if format_type == 'pdf':
        try:
            import io
            from xhtml2pdf import pisa
            out = io.BytesIO()
            pisa_status = pisa.CreatePDF(html_content, dest=out)
            if not pisa_status.err:
                response = HttpResponse(out.getvalue(), content_type='application/pdf')
                response['Content-Disposition'] = f'attachment; filename="Safety_{report_type}_{timezone.now().strftime("%Y%m%d")}.pdf"'
                return response
        except Exception:
            pass

    return HttpResponse(html_content)


# ==============================================================================
# COMPREHENSIVE MULTI-SHEET SYNC FROM OFFICIAL GOOGLE SHEET
# ==============================================================================

@login_required
def api_safety_sync_google_sheet_inventory(request):
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'Invalid method'}, status=405)

    try:
        import urllib.request, io, openpyxl, re, datetime

        url = 'https://docs.google.com/spreadsheets/d/1OBSkKFT5W01c-TsPMI_FqCGjaoLSZaxW39vZMW4lYy4/export?format=xlsx'
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        content = urllib.request.urlopen(req, timeout=25).read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

        admin_user = request.user

        # -------------------------------------------------------------
        # STEP 1: PARSE & SYNC ALL STORE INVENTORY ITEMS (GIA + 108)
        # -------------------------------------------------------------
        raw_items = {}

        def get_standard_category(cat_name, item_name):
            c = (cat_name or '').lower()
            n = item_name.lower()
            if 'helmet' in n or 'erc' in n or 'head' in c: return 'Head Protection'
            if 'reflector' in n or 'reflective' in n or 'jacket' in n or 'vis' in c: return 'Reflectors & Hi-Vis'
            if 'gumboot' in n: return 'Footwear'
            if 'shoe' in n or 'footwear' in c: return 'Safety Shoes'
            if 'glove' in n or 'hand' in c: return 'Hand Protection'
            if 'sleeve' in n or 'arm' in c: return 'Arm & Hand Protection'
            if 'leg' in n: return 'Leg Protection'
            if 'shoulder' in n: return 'Shoulder Protection'
            if 'apron' in n: return 'Body Protection'
            if 'harness' in n or 'fall' in n or 'body' in c: return 'Body Protection & Fall Safety'
            if 'goggle' in n or 'glass' in n or 'eye' in c: return 'Eye Protection'
            if 'face shield' in n or 'face protection' in c: return 'Face Protection'
            if 'plug' in n or 'muff' in n or 'hearing' in c or 'ear' in c: return 'Hearing Protection'
            if 'mask' in n or 'respirator' in n or 'respiratory' in c: return 'Respiratory Protection'
            if 'extinguisher' in n or 'blanket' in n or 'bucket' in n or 'fire' in c: return 'Fire Safety'
            if 'raincoat' in n or 'rain' in c: return 'Rainwear'
            if 'absorbent' in n or 'pad' in n or 'spill' in c: return 'Spill Control'
            return 'General Safety Equipment'

        def get_item_unit(item_name):
            n = item_name.lower()
            if 'shoe' in n or 'gumboot' in n or 'glove' in n or 'plug' in n or 'sleeve' in n or 'guard' in n or 'pad' in n:
                return 'Pairs'
            if 'tape' in n or 'net' in n:
                return 'Rolls'
            if 'harness' in n:
                return 'Sets'
            return 'Pcs'

        def get_item_warranty_and_cost(item_name):
            n = item_name.lower()
            if 'helmet' in n: return 10, 300.0
            if 'reflector' in n or 'jacket' in n: return 4, 260.0
            if 'tape' in n: return 12, 150.0
            if 'shoe' in n: return 6, 650.0
            if 'gumboot' in n: return 6, 500.0
            if 'welding' in n and 'glove' in n: return 2, 180.0
            if 'karam' in n and 'glove' in n: return 2, 120.0
            if 'throw' in n or 'sterile' in n or 'disp' in n: return 1, 30.0
            if 'glove' in n: return 2, 80.0
            if 'harness' in n or 'fall' in n: return 12, 1200.0
            if 'extinguisher' in n: return 12, 1200.0
            if 'raincoat' in n: return 12, 600.0
            if 'goggle' in n or 'glass' in n: return 6, 180.0
            if 'face shield' in n: return 6, 250.0
            if 'ear plug' in n: return 6, 90.0
            if 'ear muff' in n: return 12, 350.0
            if 'mask' in n: return 1, 40.0
            return 6, 150.0

        def make_code(name):
            clean = re.sub(r'[^a-zA-Z0-9]+', '-', name).strip('-').upper()
            parts = clean.split('-')
            if len(parts) >= 2:
                code_prefix = 'SAF-' + '-'.join(parts[:3])
            else:
                code_prefix = 'SAF-' + parts[0]
            return code_prefix[:18]

        for sname in ['HSE STORE INVENTROY (GIA)', 'HSE STORE INVENTORY (108)']:
            if sname not in wb.sheetnames: continue
            ws = wb[sname]
            for r in range(3, ws.max_row + 1):
                name = str(ws.cell(r, 2).value or '').strip()
                cat = str(ws.cell(r, 3).value or '').strip()
                qty = ws.cell(r, 4).value
                if name and name.lower() not in ['none', 'items', 's.no', 'sl.no']:
                    try: qty_int = int(float(qty))
                    except: qty_int = 0
                    clean_name = re.sub(r'\s+', ' ', name).strip().replace('[', '(').replace(']', ')')
                    clean_name = clean_name.title()
                    norm_key = clean_name.lower()
                    if norm_key not in raw_items:
                        raw_items[norm_key] = {
                            'name': clean_name,
                            'raw_cat': cat,
                            'stock': qty_int,
                            'location': 'GIA Site' if 'GIA' in sname else '108 Chorten'
                        }
                    else:
                        raw_items[norm_key]['stock'] += qty_int
                        raw_items[norm_key]['location'] = 'GIA & 108 Chorten'

        items_created = 0
        items_updated = 0
        for k, v in raw_items.items():
            name = v['name']
            cat = get_standard_category(v['raw_cat'], name)
            unit = get_item_unit(name)
            warranty, cost = get_item_warranty_and_cost(name)
            code = make_code(name)

            existing = SafetyStoreItem.objects.filter(name__iexact=name).first()
            if not existing:
                existing = SafetyStoreItem.objects.filter(item_code=code).first()

            if existing:
                if v['stock'] > 0:
                    existing.total_stock = max(existing.total_stock, v['stock'])
                    existing.available_stock = max(0, existing.total_stock - existing.issued_stock)
                existing.category = cat
                existing.unit = unit
                if not existing.warranty_months: existing.warranty_months = warranty
                if not existing.fine_amount: existing.fine_amount = Decimal(str(cost))
                existing.notes = f"Official Store Stock &bull; Location: {v['location']}"
                existing.save()
                items_updated += 1
            else:
                SafetyStoreItem.objects.create(
                    name=name,
                    item_code=code,
                    category=cat,
                    unit=unit,
                    specification=f"HSE Certified Equipment ({cat})",
                    total_stock=v['stock'],
                    available_stock=v['stock'],
                    warranty_months=warranty,
                    fine_amount=Decimal(str(cost)),
                    minimum_alert_level=10,
                    created_by=admin_user,
                    notes=f"Synced from Google Sheet &bull; Location: {v['location']}"
                )
                items_created += 1

        # -------------------------------------------------------------
        # STEP 2: PARSE & SYNC EMPLOYEE DEDUCTIONS & ISSUANCES
        # -------------------------------------------------------------
        sheet12_map = {}
        if 'Sheet12' in wb.sheetnames:
            ws12 = wb['Sheet12']
            for r in range(2, ws12.max_row + 1):
                emp_id = str(ws12.cell(r, 6).value or '').strip()
                emp_name = str(ws12.cell(r, 4).value or ws12.cell(r, 2).value or '').strip()
                ppe = str(ws12.cell(r, 8).value or '').strip().upper()
                cost = ws12.cell(r, 9).value
                rem_months = ws12.cell(r, 14).value
                fine_amt = ws12.cell(r, 15).value
                date_str = str(ws12.cell(r, 5).value or '').strip()
                if emp_id:
                    sheet12_map[emp_id.upper()] = {
                        'name': emp_name, 'ppe': ppe, 'cost': cost,
                        'rem': rem_months, 'fine': fine_amt, 'date': date_str
                    }

        def parse_sheet_date(date_val):
            if isinstance(date_val, (datetime.datetime, datetime.date)):
                return date_val if isinstance(date_val, datetime.date) else date_val.date()
            s = str(date_val or '').strip()
            if not s or s.startswith('#'):
                return datetime.date(2026, 8, 1)
            m = re.match(r'(\d{1,4})[.-/](\d{1,2})[.-/](\d{1,4})', s)
            if m:
                p1, p2, p3 = m.group(1), m.group(2), m.group(3)
                if len(p1) == 4:
                    try: return datetime.date(int(p1), int(p2), int(p3))
                    except: pass
                else:
                    yr = int(p3)
                    if yr < 100: yr += 2000
                    try: return datetime.date(yr, int(p1), int(p2))
                    except:
                        try: return datetime.date(yr, int(p2), int(p1))
                        except: pass
            return datetime.date(2026, 8, 1)

        def parse_remarks_for_months(remarks_str):
            used = 0
            rem = 0
            s = str(remarks_str or '').upper()
            m_comp = re.search(r'COMPLETED\s+(\d+)\s+MONTH', s)
            if m_comp: used = int(m_comp.group(1))
            m_left = re.search(r'LEFT\s+(\d+)\s+MONTH', s)
            if m_left: rem = int(m_left.group(1))
            if used == 0 and 'MONTH' in s:
                m_any = re.search(r'(\d+)\s+MONTH', s)
                if m_any: used = int(m_any.group(1))
            return used, rem

        synced_fines = 0
        synced_issues = 0

        for sname, dept_prefix in [
            ('Company Employee for Money Dedu', 'RVJV Company Employee'),
            ('Hiring Employee for Money Deduc', 'Buddha Dev Hiring Labour')
        ]:
            if sname not in wb.sheetnames: continue
            ws = wb[sname]
            start_row = 3 if 'Company' in sname else 4
            for r in range(start_row, ws.max_row + 1):
                name = str(ws.cell(r, 2).value or '').strip()
                raw_date = ws.cell(r, 3).value
                emp_id = str(ws.cell(r, 4).value or '').strip()
                desig = str(ws.cell(r, 5).value or '').strip()
                reflector_val = str(ws.cell(r, 6).value or '').strip()
                shoes_val = str(ws.cell(r, 7).value or '').strip()
                helmet_val = str(ws.cell(r, 8).value or '').strip()
                remarks = str(ws.cell(r, 9).value or '').strip()

                if not name or name.lower() in ['none', 'name', 'sl.no', 's.no']:
                    continue

                rep_date = parse_sheet_date(raw_date)
                used_m, rem_m = parse_remarks_for_months(remarks)

                emp = None
                if emp_id and emp_id.lower() not in ['none', '']:
                    emp = Employee.objects.filter(emp_id__iexact=emp_id).first()
                if not emp:
                    emp = Employee.objects.filter(name__iexact=name).first()
                if not emp:
                    clean_id = emp_id if (emp_id and len(emp_id) > 2) else f"HIR-{r:04d}"
                    emp = Employee.objects.create(
                        emp_id=clean_id,
                        name=name.title(),
                        department=f"{dept_prefix} ({desig or 'Site Worker'})",
                        status='ACTIVE'
                    )

                target_items = []
                if shoes_val and shoes_val.lower() not in ['none', '']:
                    size_m = re.search(r'\((\d+)\)', shoes_val)
                    size_str = size_m.group(1) if size_m else ('8' if '8' in shoes_val else '8')
                    shoe_item = SafetyStoreItem.objects.filter(name__icontains=f"({size_str})").filter(name__icontains="shoe").first()
                    if not shoe_item:
                        shoe_item = SafetyStoreItem.objects.filter(category='Safety Shoes').first()
                    if shoe_item:
                        target_items.append((shoe_item, f"Size {size_str}", 650.0, 6))

                if reflector_val and reflector_val.lower() not in ['none', '']:
                    ref_item = SafetyStoreItem.objects.filter(category__icontains='Reflector').first()
                    if ref_item:
                        target_items.append((ref_item, "Standard", 260.0, 4))

                if helmet_val and helmet_val.lower() not in ['none', '']:
                    hlm_item = SafetyStoreItem.objects.filter(category__icontains='Head').first()
                    if hlm_item:
                        target_items.append((hlm_item, "Standard", 300.0, 10))

                if not target_items:
                    fallback = SafetyStoreItem.objects.filter(category='Safety Shoes').first()
                    if fallback:
                        target_items.append((fallback, "Standard", 650.0, 6))

                s12_info = sheet12_map.get(emp_id.upper()) if emp_id else None

                for item, size_spec, full_cost, total_life_m in target_items:
                    if s12_info and s12_info.get('fine') is not None:
                        try: fine_amount = Decimal(str(float(s12_info['fine'])))
                        except: fine_amount = Decimal('100.00')
                        if s12_info.get('rem'):
                            try: rem_m = int(float(s12_info['rem']))
                            except: pass
                    else:
                        if rem_m > 0 and total_life_m > 0:
                            calculated = (Decimal(str(full_cost)) / Decimal(str(total_life_m))) * Decimal(str(rem_m))
                            fine_amount = round(calculated, 2)
                        else:
                            fine_amount = Decimal('100.00')

                    issue_date = rep_date - datetime.timedelta(days=max(30, used_m * 30))
                    expiry_date = issue_date + datetime.timedelta(days=total_life_m * 30)

                    issue_obj, iss_created = SafetyEquipmentIssue.objects.get_or_create(
                        employee=emp,
                        item=item,
                        issue_date=issue_date,
                        defaults={
                            'quantity': 1,
                            'size_specification': size_spec,
                            'warranty_expiry_date': expiry_date,
                            'status': 'REPLACED_PREMATURE' if rem_m > 0 else 'ACTIVE',
                            'issued_by': admin_user,
                            'remarks': f"Imported from HSE Sheet &bull; {remarks}"
                        }
                    )
                    if iss_created: synced_issues += 1

                    fine_obj, fine_created = SafetyReplacementAndFine.objects.get_or_create(
                        employee=emp,
                        item=item,
                        replacement_date=rep_date,
                        defaults={
                            'original_issue': issue_obj,
                            'reason': 'DAMAGED_PREMATURE',
                            'old_item_condition': f"Used: {used_m} mos, Returned premature: {rem_m} mos remaining",
                            'is_premature': (rem_m > 0),
                            'days_used': used_m * 30,
                            'warranty_days': total_life_m * 30,
                            'months_used': used_m,
                            'remaining_months': rem_m,
                            'full_cost': Decimal(str(full_cost)),
                            'fine_amount': fine_amount,
                            'fine_status': 'PENDING_APPROVAL',
                            'processed_by': admin_user
                        }
                    )
                    if fine_created: synced_fines += 1

        msg = (
            f"Successfully Synced All Sheets from Google Sheet!\n"
            f"• Store Catalog: {items_created} Added, {items_updated} Updated (Total: {SafetyStoreItem.objects.count()} items)\n"
            f"• Deductions & Replacements: {synced_fines} Fine Records, {synced_issues} Issue Records synced."
        )
        return JsonResponse({
            'status': 'SUCCESS',
            'created_count': items_created,
            'updated_count': items_updated,
            'synced_fines': synced_fines,
            'synced_issues': synced_issues,
            'msg': msg
        })

    except Exception as e:
        return JsonResponse({'status': 'ERROR', 'msg': f"Sync failed: {str(e)}"})


# ==============================================================================
# EXCEL FILE UPLOAD & SYNC API
# ==============================================================================

@login_required
def api_safety_excel_upload(request):
    """
    Upload safety_office_data.xlsx and sync:
      1. Store Inventory — from the latest monthly stock sheet
      2. Employee Issuances — from the Employee sheet (only employees in that sheet)
    """
    if not has_safety_access(request.user):
        return JsonResponse({'status': 'ERROR', 'msg': 'Permission denied.'}, status=403)

    if request.method != 'POST':
        return JsonResponse({'status': 'ERROR', 'msg': 'POST required.'}, status=405)

    uploaded_file = request.FILES.get('file')
    if not uploaded_file:
        return JsonResponse({'status': 'ERROR', 'msg': 'No file uploaded.'})

    if not uploaded_file.name.endswith('.xlsx'):
        return JsonResponse({'status': 'ERROR', 'msg': 'Only .xlsx files are supported.'})

    try:
        import io
        import re
        import uuid as _uuid

        wb = openpyxl.load_workbook(io.BytesIO(uploaded_file.read()), data_only=True)
        today = timezone.now().date()

        results = {
            'items_created': 0,
            'items_updated': 0,
            'issues_created': 0,
            'issues_skipped': 0,
            'emp_not_found': 0,
            'errors': [],
        }

        # ── Helper: normalise a name for matching ──────────────────────────────
        def _norm(s):
            return re.sub(r'\s+', ' ', str(s or '').strip().lower())

        # ── Helper: guess category from item name ──────────────────────────────
        def _guess_category(name):
            n = name.lower()
            if any(k in n for k in ['helmet', 'hard hat']): return 'HELMET'
            if any(k in n for k in ['shoe', 'boot', 'gumboot']): return 'SHOES'
            if any(k in n for k in ['vest', 'jacket', 'reflector', 'hi viz', 'hiviz', 'reflective']): return 'VEST'
            if any(k in n for k in ['glove', 'sleeve', 'apron']): return 'GLOVES'
            if any(k in n for k in ['harness', 'lanyard']): return 'HARNESS'
            if any(k in n for k in ['goggle', 'glass', 'shield', 'visor']): return 'GOGGLES'
            if any(k in n for k in ['ear plug', 'ear muff', 'earplug', 'earmuff']): return 'EARPLUG'
            if any(k in n for k in ['mask', 'respirator', 'gas filter']): return 'RESPIRATOR'
            return 'OTHER'

        def _guess_unit(name):
            n = name.lower()
            if any(k in n for k in ['glove', 'boot', 'shoe', 'gumboot']): return 'Pairs'
            if any(k in n for k in ['tape', 'net']): return 'Rolls'
            if any(k in n for k in ['extinguisher', 'machine', 'light', 'torch']): return 'Units'
            return 'Pcs'

        # ──────────────────────────────────────────────────────────────────────
        # PART 1: SYNC STORE INVENTORY from the latest monthly stock sheet
        # ──────────────────────────────────────────────────────────────────────
        STOCK_SHEET_PRIORITY = [
            'JULY STOCK', 'JUNE STOCK', 'may stocks', 'APRIL STOCKS',
            'MARCH STOCKS', 'FEBURARY STOCKS'
        ]

        stock_sheet = None
        for sheet_name in STOCK_SHEET_PRIORITY:
            if sheet_name in wb.sheetnames:
                ws_test = wb[sheet_name]
                # Check if it has actual numeric stock data (col C = STOCKS AVAILABLE)
                has_data = False
                for row in ws_test.iter_rows(min_row=2, max_row=10, values_only=True):
                    if row[1] and row[2] is not None:
                        try:
                            val = float(str(row[2]))
                            has_data = True
                            break
                        except (ValueError, TypeError):
                            continue
                if has_data:
                    stock_sheet = ws_test
                    break

        if stock_sheet is not None:
            # Build lookup: normalised_name → SafetyStoreItem
            existing_items = {_norm(item.name): item for item in SafetyStoreItem.objects.all()}

            for row in stock_sheet.iter_rows(min_row=2, values_only=True):
                item_name = row[1] if len(row) > 1 else None
                stock_val = row[2] if len(row) > 2 else None

                if not item_name:
                    continue

                item_name_str = str(item_name).strip()
                if not item_name_str or item_name_str.upper() in ('ITEMS NAME', 'SL.NO'):
                    continue

                # Parse stock — skip formula strings
                try:
                    stock_qty = int(float(str(stock_val)))
                    if stock_qty < 0:
                        stock_qty = 0
                except (ValueError, TypeError):
                    continue  # skip rows with formula or blank stock

                norm_name = _norm(item_name_str)

                if norm_name in existing_items:
                    # Update existing item
                    store_item = existing_items[norm_name]
                    store_item.total_stock = stock_qty
                    store_item.available_stock = stock_qty
                    store_item.save(update_fields=['total_stock', 'available_stock', 'updated_at'])
                    results['items_updated'] += 1
                else:
                    # Create new item
                    try:
                        auto_code = f"SAF-{_uuid.uuid4().hex[:6].upper()}"
                        while SafetyStoreItem.objects.filter(item_code=auto_code).exists():
                            auto_code = f"SAF-{_uuid.uuid4().hex[:6].upper()}"

                        new_item = SafetyStoreItem.objects.create(
                            name=item_name_str.title(),
                            item_code=auto_code,
                            category=_guess_category(item_name_str),
                            unit=_guess_unit(item_name_str),
                            total_stock=stock_qty,
                            available_stock=stock_qty,
                            minimum_alert_level=5,
                            warranty_months=6,
                            fine_amount=0,
                            created_by=request.user,
                        )
                        existing_items[norm_name] = new_item
                        results['items_created'] += 1
                    except Exception as e:
                        results['errors'].append(f"Item '{item_name_str}': {str(e)}")

        # ──────────────────────────────────────────────────────────────────────
        # PART 2: SYNC EMPLOYEE ISSUANCES from Employee sheet
        # ──────────────────────────────────────────────────────────────────────
        if 'Employee' not in wb.sheetnames:
            results['errors'].append("Employee sheet not found in uploaded file.")
        else:
            emp_ws = wb['Employee']
            all_headers = [str(c.value).strip() if c.value else '' for c in emp_ws[1]]

            # Equipment columns start at index 7 (after Sr, ID, Full Name, Designation, Mobile, Status, CID)
            EQUIP_START_INDEX = 7
            equip_headers = all_headers[EQUIP_START_INDEX:]

            # Build store item lookup from DB (after Part 1 may have added new items)
            store_lookup = {}
            for si in SafetyStoreItem.objects.all():
                store_lookup[_norm(si.name)] = si

            # Build equipment name → SafetyStoreItem mapping from column headers
            equip_to_item = {}
            # Mapping from Excel column name → likely item name patterns
            EQUIP_NAME_MAP = {
                'helmet': ['helmet', 'hard hat'],
                'hi viz jacket': ['hi viz', 'hiviz', 'reflector jacket', 'green jacket', 'jacket reflector'],
                'safety boots': ['safety shoe', 'boot', 'safety boot'],
                'rubber hand gloves': ['rubber hand glove', 'blue long hand glove', 'hand glove'],
                'ear plugs': ['ear plug'],
                'hand gloves': ['karam hand glove', 'hand glove', 'glove'],
                'safety goggles': ['safety goggle', 'goggle'],
                'sign boards': ['sign board'],
                'full body harness': ['full body harness', 'harness'],
                'reflective tapes': ['reflective tape', 'tape'],
                'barrification tape': ['barrication tape', 'barrification tape'],
                'cone': ['cone'],
                'face mask': ['face mask', 'mask'],
                'alchol testing machine': ['alcohol testing machine', 'alchol testing'],
                'fire extinguisher': ['fire extinguisher'],
                'shoulder pads': ['shoulder pad'],
                'green nets': ['green net', 'net'],
                'bucket': ['fire bucket', 'bucket'],
                'torch': ['torch', 'search light', 'led light'],
                'whistle': ['whistle'],
                'body harness': ['body harness', 'harness'],
                'traffic light': ['traffic light'],
                'hand sleeve': ['welding hand sleeve', 'hand sleeve'],
                'weldingglove': ['welding hand glove', 'welding glove'],
                'leg guard': ['welding leg guard', 'leg guard'],
                'apron': ['apron'],
                'face shield': ['face shield'],
                'raincoat': ['raincoat', 'waterproof', 'water proof'],
                'welding helmet': ['welding helmet'],
                'black glass': ['black glass', 'black goggle'],
                'head screen': ['head screen'],
                'black goggle': ['black goggle', 'black glass'],
                'ear muffin': ['ear muff', 'karam ep'],
                'white glass': ['white glass'],
                'use throw glove': ['use and throw', 'nitrile', 'throw glove'],
            }

            for col_idx, col_name in enumerate(equip_headers):
                if not col_name:
                    continue
                col_norm = _norm(col_name)

                # Try direct match first
                if col_norm in store_lookup:
                    equip_to_item[col_idx] = store_lookup[col_norm]
                    continue

                # Try mapped patterns
                patterns = EQUIP_NAME_MAP.get(col_norm, [col_norm])
                matched = None
                for pat in patterns:
                    for store_norm, store_item in store_lookup.items():
                        if pat in store_norm or store_norm in pat:
                            matched = store_item
                            break
                    if matched:
                        break

                if matched:
                    equip_to_item[col_idx] = matched

            # Build emp_id lookup
            emp_lookup = {emp.emp_id.strip(): emp for emp in Employee.objects.all() if emp.emp_id}

            # Process each employee row
            for row in emp_ws.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None:
                    continue

                emp_id_raw = str(row[1]).strip() if row[1] else ''
                status_raw = str(row[5]).strip().lower() if row[5] else ''

                if not emp_id_raw:
                    continue

                # Find employee
                emp_obj = emp_lookup.get(emp_id_raw)
                if not emp_obj:
                    results['emp_not_found'] += 1
                    continue

                # Process each equipment column
                equip_data = row[EQUIP_START_INDEX:]
                for col_idx, cell_val in enumerate(equip_data):
                    if cell_val is None or cell_val == '':
                        continue

                    # Parse quantity — handle "1[23/6]" style values
                    qty = 0
                    cell_str = str(cell_val).strip()
                    qty_match = re.match(r'(\d+)', cell_str)
                    if qty_match:
                        qty = int(qty_match.group(1))

                    if qty <= 0:
                        continue

                    store_item = equip_to_item.get(col_idx)
                    if not store_item:
                        continue

                    # Check if active issue already exists for this employee + item
                    already_exists = SafetyEquipmentIssue.objects.filter(
                        employee=emp_obj,
                        item=store_item,
                        status='ACTIVE'
                    ).exists()

                    if already_exists:
                        results['issues_skipped'] += 1
                        continue

                    try:
                        SafetyEquipmentIssue.objects.create(
                            employee=emp_obj,
                            item=store_item,
                            quantity=qty,
                            issue_date=today,
                            status='ACTIVE',
                            issued_by=request.user,
                            remarks=f'Imported from safety_office_data.xlsx upload'
                        )
                        results['issues_created'] += 1
                    except Exception as e:
                        results['errors'].append(f"Issue for {emp_id_raw} / {store_item.name}: {str(e)}")

        # ── Build summary message ──────────────────────────────────────────────
        msg = (
            f"✅ Excel sync complete!\n"
            f"📦 Store Inventory: {results['items_created']} new items added, {results['items_updated']} items updated.\n"
            f"👷 Issuances: {results['issues_created']} new records created, {results['issues_skipped']} already existed (skipped).\n"
            f"⚠️ Employees not found in system: {results['emp_not_found']}."
        )
        if results['errors']:
            msg += f"\n❌ Errors ({len(results['errors'])}): " + "; ".join(results['errors'][:5])

        return JsonResponse({
            'status': 'SUCCESS',
            'msg': msg,
            'items_created': results['items_created'],
            'items_updated': results['items_updated'],
            'issues_created': results['issues_created'],
            'issues_skipped': results['issues_skipped'],
            'emp_not_found': results['emp_not_found'],
            'errors': results['errors'][:10],
        })

    except Exception as e:
        import traceback
        return JsonResponse({'status': 'ERROR', 'msg': f"Upload failed: {str(e)}", 'trace': traceback.format_exc()})


# ─────────────────────────────────────────────────────────────────────────────
# MANPOWER DASHBOARD
# ─────────────────────────────────────────────────────────────────────────────
@login_required
def manpower_dashboard_view(request):
    from .models import LabourRecord, RVJVEmployee, HiredOperator, ContractorWorkerPPE
    from django.db.models import Q

    active_tab = request.GET.get('tab', 'labour')

    # ── TAB A: Labour Register ────────────────────────────────────────────────
    labour_qs = LabourRecord.objects.all().order_by('name')
    lab_status_filter = request.GET.get('lab_status', '')
    lab_nat_filter    = request.GET.get('lab_nat', '')
    lab_search        = request.GET.get('lab_q', '').strip()
    if lab_status_filter:
        labour_qs = labour_qs.filter(status=lab_status_filter)
    if lab_nat_filter:
        labour_qs = labour_qs.filter(nationality=lab_nat_filter)
    if lab_search:
        labour_qs = labour_qs.filter(
            Q(name__icontains=lab_search) | Q(labour_id__icontains=lab_search) |
            Q(cid_number__icontains=lab_search)
        )

    # ── TAB B: RVJV Employees ─────────────────────────────────────────────────
    rvjv_qs = RVJVEmployee.objects.all().order_by('name')
    rvjv_dept_filter   = request.GET.get('rvjv_dept', '')
    rvjv_status_filter = request.GET.get('rvjv_status', '')
    rvjv_nat_filter    = request.GET.get('rvjv_nat', '')
    rvjv_search        = request.GET.get('rvjv_q', '').strip()
    if rvjv_dept_filter:
        rvjv_qs = rvjv_qs.filter(department=rvjv_dept_filter)
    if rvjv_status_filter:
        rvjv_qs = rvjv_qs.filter(status=rvjv_status_filter)
    if rvjv_nat_filter:
        rvjv_qs = rvjv_qs.filter(nationality=rvjv_nat_filter)
    if rvjv_search:
        rvjv_qs = rvjv_qs.filter(
            Q(name__icontains=rvjv_search) | Q(emp_id__icontains=rvjv_search) |
            Q(designation__icontains=rvjv_search) | Q(cid_number__icontains=rvjv_search)
        )

    # ── TAB C: Hired Operators ────────────────────────────────────────────────
    hired_qs = HiredOperator.objects.all().order_by('name')
    hired_agent_filter = request.GET.get('hired_agent', '')
    hired_type_filter  = request.GET.get('hired_type', '')
    hired_nat_filter   = request.GET.get('hired_nat', '')
    hired_search       = request.GET.get('hired_q', '').strip()
    if hired_agent_filter:
        hired_qs = hired_qs.filter(hire_agent=hired_agent_filter)
    if hired_type_filter:
        hired_qs = hired_qs.filter(hire_type=hired_type_filter)
    if hired_nat_filter:
        hired_qs = hired_qs.filter(nationality=hired_nat_filter)
    if hired_search:
        hired_qs = hired_qs.filter(
            Q(name__icontains=hired_search) | Q(operator_id__icontains=hired_search) |
            Q(vehicle_no__icontains=hired_search) | Q(designation__icontains=hired_search)
        )

    # ── TAB D: Contractor PPE ─────────────────────────────────────────────────
    ppe_qs = ContractorWorkerPPE.objects.all().order_by('contractor_name', 'name')
    ppe_contractor_filter = request.GET.get('ppe_contractor', '')
    ppe_desig_filter      = request.GET.get('ppe_desig', '')
    ppe_search            = request.GET.get('ppe_q', '').strip()
    if ppe_contractor_filter:
        ppe_qs = ppe_qs.filter(contractor_name=ppe_contractor_filter)
    if ppe_desig_filter:
        ppe_qs = ppe_qs.filter(designation__icontains=ppe_desig_filter)
    if ppe_search:
        ppe_qs = ppe_qs.filter(
            Q(name__icontains=ppe_search) | Q(worker_id__icontains=ppe_search) |
            Q(contractor_name__icontains=ppe_search)
        )

    # ── Filter dropdown options ────────────────────────────────────────────────
    all_rvjv_depts  = RVJVEmployee.objects.values_list('department', flat=True).exclude(department=None).distinct().order_by('department')
    all_hire_agents = HiredOperator.objects.values_list('hire_agent', flat=True).exclude(hire_agent=None).distinct().order_by('hire_agent')
    all_contractors = ContractorWorkerPPE.objects.values_list('contractor_name', flat=True).distinct().order_by('contractor_name')

    # ── KPI Counts (always unfiltered) ────────────────────────────────────────
    labour_total     = LabourRecord.objects.count()
    labour_active    = LabourRecord.objects.filter(status='Active').count()
    labour_bhutanese = LabourRecord.objects.filter(nationality='Bhutanese').count()
    labour_indian    = LabourRecord.objects.filter(nationality='Indian').count()

    rvjv_total   = RVJVEmployee.objects.count()
    rvjv_active  = RVJVEmployee.objects.filter(status='Active').count()
    rvjv_inactive = RVJVEmployee.objects.filter(status='Inactive').count()
    rvjv_left    = RVJVEmployee.objects.filter(status='Left').count()

    hired_total    = HiredOperator.objects.count()
    hired_driver   = HiredOperator.objects.filter(hire_type='Driver/Operator').count()
    hired_mechanic = HiredOperator.objects.filter(hire_type='Mechanic/Supervisor').count()

    ppe_total       = ContractorWorkerPPE.objects.count()
    ppe_contractors = ContractorWorkerPPE.objects.values('contractor_name').distinct().count()

    context = {
        'active_tab': active_tab,
        # Querysets
        'labour_list': labour_qs,
        'rvjv_list':   rvjv_qs,
        'hired_list':  hired_qs,
        'ppe_list':    ppe_qs,
        # Filter values
        'lab_status_filter': lab_status_filter,
        'lab_nat_filter':    lab_nat_filter,
        'lab_search':        lab_search,
        'rvjv_dept_filter':   rvjv_dept_filter,
        'rvjv_status_filter': rvjv_status_filter,
        'rvjv_nat_filter':    rvjv_nat_filter,
        'rvjv_search':        rvjv_search,
        'hired_agent_filter': hired_agent_filter,
        'hired_type_filter':  hired_type_filter,
        'hired_nat_filter':   hired_nat_filter,
        'hired_search':       hired_search,
        'ppe_contractor_filter': ppe_contractor_filter,
        'ppe_desig_filter':      ppe_desig_filter,
        'ppe_search':            ppe_search,
        # Dropdown options
        'all_rvjv_depts':  all_rvjv_depts,
        'all_hire_agents': all_hire_agents,
        'all_contractors': all_contractors,
        # KPIs
        'labour_total': labour_total, 'labour_active': labour_active,
        'labour_bhutanese': labour_bhutanese, 'labour_indian': labour_indian,
        'rvjv_total': rvjv_total, 'rvjv_active': rvjv_active,
        'rvjv_inactive': rvjv_inactive, 'rvjv_left': rvjv_left,
        'hired_total': hired_total, 'hired_driver': hired_driver,
        'hired_mechanic': hired_mechanic,
        'ppe_total': ppe_total, 'ppe_contractors': ppe_contractors,
    }
    return render(request, 'manpower_dashboard.html', context)
