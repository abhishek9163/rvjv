from django.contrib import admin
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

@admin.action(description='Export Selected to Excel (.xlsx)')
def export_as_excel(modeladmin, request, queryset):
    meta = modeladmin.model._meta
    field_names = [field.name for field in meta.fields]
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename={meta.model_name}s.xlsx'
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = meta.model_name.capitalize()[:30]
    
    # Headers styling matching corporate theme
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(left=Side(style='thin', color='DDDDDD'),
                         right=Side(style='thin', color='DDDDDD'),
                         top=Side(style='thin', color='DDDDDD'),
                         bottom=Side(style='thin', color='DDDDDD'))
    
    # Write headers
    headers = [field.verbose_name.upper() if hasattr(field, 'verbose_name') else field.name.upper() for field in meta.fields]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    # Write data rows
    for obj in queryset:
        row = []
        for field in field_names:
            val = getattr(obj, field)
            if val is None:
                val = ""
            elif hasattr(val, 'strftime'):
                val = val.strftime('%Y-%m-%d %H:%M:%S') if hasattr(val, 'hour') else val.strftime('%Y-%m-%d')
            elif hasattr(val, 'url'):
                val = request.build_absolute_uri(val.url)
            else:
                val = str(val)
            row.append(val)
        ws.append(row)
        
    # Autofit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        for cell in col:
            cell.border = thin_border
            
    wb.save(response)
    return response

@admin.action(description='♻️ Restore Selected Deleted Entries')
def restore_selected(modeladmin, request, queryset):
    restored_count = 0
    for obj in queryset:
        if getattr(obj, 'is_deleted', False):
            if obj.__class__.__name__ == 'SparePartTransaction':
                part = obj.part
                if obj.transaction_type == 'IN':
                    part.current_stock += obj.quantity
                else:
                    part.current_stock -= obj.quantity
                part.save()
            obj.is_deleted = False
            obj.deleted_at = None
            obj.deleted_by = None
            obj.save()
            restored_count += 1
    
    if restored_count > 0:
        modeladmin.message_user(request, f"Successfully restored {restored_count} entries.")
    else:
        modeladmin.message_user(request, "No deleted entries were selected or restored.", level='warning')

# Custom Admin supporting advanced top-bar export buttons
class ExportableAdmin(admin.ModelAdmin):
    change_list_template = 'admin/custom_change_list.html'
    excel_export_url = True
    pdf_export_url = True
    
    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context['excel_export_url'] = self.excel_export_url
        extra_context['pdf_export_url'] = self.pdf_export_url
        # Pass model meta info as safe context (templates can't access _meta directly)
        extra_context['model_fields'] = [
            {'name': f.name, 'verbose_name': str(f.verbose_name).title()}
            for f in self.model._meta.fields
        ]
        extra_context['model_name'] = self.model._meta.model_name
        return super().changelist_view(request, extra_context=extra_context)
        
    def get_urls(self):
        from django.urls import path
        urls = super().get_urls()
        custom_urls = [
            path('export-excel-admin/', self.admin_site.admin_view(self.export_excel_admin_view), name=f'{self.model._meta.model_name}_export_excel_admin'),
            path('export-pdf-admin/', self.admin_site.admin_view(self.export_pdf_admin_view), name=f'{self.model._meta.model_name}_export_pdf_admin'),
        ]
        return custom_urls + urls

    def export_excel_admin_view(self, request):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        # Get the filtered queryset
        cl = self.get_changelist_instance(request)
        queryset = cl.get_queryset(request)
        
        meta = self.model._meta
        
        # Read column parameter
        columns_param = request.GET.get('columns')
        if columns_param:
            field_names = [f.strip() for f in columns_param.split(',') if f.strip()]
        else:
            field_names = [field.name for field in meta.fields]
            
        # Optional vendor filtering
        vendor_name = request.GET.get('vendor_name')
        if vendor_name:
            if hasattr(self.model, 'supplier'):
                queryset = queryset.filter(supplier=vendor_name)
            elif hasattr(self.model, 'owner_name'):
                queryset = queryset.filter(owner_name=vendor_name)
                
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename={meta.model_name}s_filtered.xlsx'
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = meta.model_name.capitalize()[:30]
        
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(left=Side(style='thin', color='DDDDDD'),
                             right=Side(style='thin', color='DDDDDD'),
                             top=Side(style='thin', color='DDDDDD'),
                             bottom=Side(style='thin', color='DDDDDD'))
        
        # Build headers
        headers = []
        for fn in field_names:
            try:
                field = meta.get_field(fn)
                headers.append(field.verbose_name.upper() if hasattr(field, 'verbose_name') else fn.upper())
            except:
                headers.append(fn.upper())
                
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for obj in queryset:
            row = []
            for field in field_names:
                try:
                    val = getattr(obj, field)
                    if val is None:
                        val = ""
                    elif hasattr(val, 'strftime'):
                        val = val.strftime('%Y-%m-%d %H:%M:%S') if hasattr(val, 'hour') else val.strftime('%Y-%m-%d')
                    elif hasattr(val, 'url'):
                        val = request.build_absolute_uri(val.url)
                    elif hasattr(val, 'username'):  # For ForeignKeys e.g. User or related object
                        val = str(val)
                    else:
                        val = str(val)
                except:
                    val = ""
                row.append(val)
            ws.append(row)
            
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
            for cell in col:
                cell.border = thin_border
                
        wb.save(response)
        return response

    def export_pdf_admin_view(self, request):
        from django.utils import timezone
        
        cl = self.get_changelist_instance(request)
        queryset = cl.get_queryset(request)
        
        meta = self.model._meta
        
        columns_param = request.GET.get('columns')
        if columns_param:
            field_names = [f.strip() for f in columns_param.split(',') if f.strip()]
        else:
            field_names = [field.name for field in meta.fields]
            
        # Optional vendor filtering
        vendor_name = request.GET.get('vendor_name')
        if vendor_name:
            if hasattr(self.model, 'supplier'):
                queryset = queryset.filter(supplier=vendor_name)
            elif hasattr(self.model, 'owner_name'):
                queryset = queryset.filter(owner_name=vendor_name)
                
        headers = []
        for fn in field_names:
            try:
                field = meta.get_field(fn)
                headers.append(field.verbose_name.upper() if hasattr(field, 'verbose_name') else fn.upper())
            except:
                headers.append(fn.upper())
        
        html = f"""
        <html>
        <head>
            <title>{meta.verbose_name_plural.title()}</title>
            <style>
                body {{ font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif; margin: 20px; color: #333; }}
                h2 {{ text-align: center; color: #1e3a8a; margin-bottom: 20px; text-transform: uppercase; font-size: 16px; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 10px; }}
                th, td {{ border: 1px solid #ddd; padding: 6px 8px; text-align: left; }}
                th {{ background-color: #1e3a8a; color: white; font-weight: bold; text-transform: uppercase; }}
                tr:nth-child(even) {{ background-color: #f8fafc; }}
                .footer {{ text-align: right; margin-top: 20px; font-size: 9px; color: #666; }}
                @media print {{
                    @page {{ size: landscape; margin: 1cm; }}
                    button {{ display: none; }}
                }}
            </style>
        </head>
        <body>
            <h2>{meta.verbose_name_plural.title()} Report</h2>
            <table>
                <thead>
                    <tr>
        """
        
        for h in headers:
            html += f"<th>{h}</th>"
        html += "</tr></thead><tbody>"
        
        for obj in queryset:
            html += "<tr>"
            for field in field_names:
                try:
                    val = getattr(obj, field)
                    if val is None:
                        val = ""
                    elif hasattr(val, 'strftime'):
                        val = val.strftime('%Y-%m-%d %H:%M:%S') if hasattr(val, 'hour') else val.strftime('%Y-%m-%d')
                    elif hasattr(val, 'url'):
                        val = request.build_absolute_uri(val.url)
                    else:
                        val = str(val)
                except:
                    val = ""
                html += f"<td>{val}</td>"
            html += "</tr>"
            
        html += f"""
                </tbody>
            </table>
            <div class="footer">Generated on: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}</div>
            <script>window.onload = function() {{ window.print(); }}</script>
        </body>
        </html>
        """
        return HttpResponse(html)

from .models import Driver, VehicleModel, FleetVehicle, VehicleMovement, ServiceLog, RepairLog, LubricationLog, TyreLog, HiredVehicle, SparePart, SparePartTransaction

class FleetVehicleAdmin(ExportableAdmin):
    actions = [export_as_excel]
    excel_export_url = True
    pdf_export_url = True
    list_display = ['dno', 'regn', 'model_name', 'department', 'is_active']
    list_filter = ['is_active', 'department']
    search_fields = ['dno', 'regn', 'model_name']

class DriverAdmin(admin.ModelAdmin):
    actions = [export_as_excel]
    list_display = ['name', 'phone', 'license_no', 'is_active']
    list_filter = ['is_active']
    search_fields = ['name', 'phone', 'license_no']

class RepairLogAdmin(admin.ModelAdmin):
    actions = [export_as_excel]
    list_display = ['vehicle', 'in_date', 'out_date', 'status']
    search_fields = ['vehicle__dno', 'vehicle__regn', 'mechanic']

class VehicleMovementAdmin(ExportableAdmin):
    actions = [export_as_excel, restore_selected]
    excel_export_url = True
    pdf_export_url = False  # Movements has no PDF export template, so we disable it or use actions
    list_display = ['vehicle_number', 'driver_name', 'movement_date', 'movement_time', 'shift', 'destination', 'is_deleted']
    list_filter = ['shift', 'movement_date', 'is_deleted']
    search_fields = ['vehicle_number', 'driver_name', 'destination']

class LubricationLogAdmin(ExportableAdmin):
    actions = [export_as_excel, restore_selected]
    excel_export_url = True
    pdf_export_url = True
    list_display = ['vehicle', 'date', 'oil_type', 'qty', 'unit', 'total_amount', 'is_deleted']
    list_filter = ['oil_type', 'date', 'is_deleted']
    search_fields = ['vehicle__regn', 'oil_type']

class TyreLogAdmin(ExportableAdmin):
    actions = [export_as_excel, restore_selected]
    excel_export_url = True
    pdf_export_url = True
    list_display = ['vehicle', 'date', 'location', 'work_order_no', 'vendor', 'punctures', 'total_amount', 'is_deleted']
    list_filter = ['date', 'vendor', 'is_deleted']
    search_fields = ['vehicle__regn', 'vendor', 'location']

class HiredVehicleAdmin(admin.ModelAdmin):
    actions = [export_as_excel]
    list_display = ['regn', 'equipment_type', 'owner_name', 'status', 'contract_status']
    list_filter = ['status', 'contract_status', 'equipment_type']
    search_fields = ['regn', 'owner_name', 'equipment_type']

class SparePartAdmin(admin.ModelAdmin):
    actions = [export_as_excel]
    list_display = ['part_number', 'part_name', 'category', 'current_stock', 'reorder_level', 'valuation_rate']
    list_filter = ['category']
    search_fields = ['part_number', 'part_name']

class SparePartTransactionAdmin(ExportableAdmin):
    actions = [export_as_excel, restore_selected]
    excel_export_url = True
    pdf_export_url = True
    list_display = ['part', 'transaction_type', 'date', 'quantity', 'rate', 'total_amount', 'is_approved', 'is_deleted']
    list_filter = ['transaction_type', 'is_approved', 'date', 'is_deleted']
    search_fields = ['part__part_number', 'part__part_name', 'supplier']

admin.site.register(VehicleModel)
admin.site.register(Driver, DriverAdmin)
admin.site.register(FleetVehicle, FleetVehicleAdmin)
admin.site.register(VehicleMovement, VehicleMovementAdmin)
admin.site.register(ServiceLog)
admin.site.register(RepairLog, RepairLogAdmin)
admin.site.register(LubricationLog, LubricationLogAdmin)
admin.site.register(TyreLog, TyreLogAdmin)
admin.site.register(HiredVehicle, HiredVehicleAdmin)
admin.site.register(SparePart, SparePartAdmin)
admin.site.register(SparePartTransaction, SparePartTransactionAdmin)
